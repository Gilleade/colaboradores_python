from __future__ import annotations

import shutil
import csv
import io
import os
import re
import sys
import unicodedata
import subprocess
import threading
import traceback
import webbrowser
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple
from difflib import SequenceMatcher

from license_validator import (
    LicenseValidationError,
    LicenseValidationResult,
    get_current_machine_id,
    save_activation_request,
    validate_license,
)

#Arquivo XLSX
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

import fitz  # PyMuPDF
import pytesseract
from PIL import Image, ImageEnhance, ImageFilter, ImageOps, ImageTk

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except Exception:  # pragma: no cover
    tk = None
    filedialog = None
    messagebox = None
    ttk = None


# ============================================================
# Configuração do aplicativo
# ============================================================

def get_app_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


APP_DIR = get_app_dir()
BASE_DIR = APP_DIR / "colaboradores_data"
BRUTO_DIR = BASE_DIR / "bruto"
PROCESSADO_DIR = BASE_DIR / "processado"
ACTIVATION_REQUEST_PATH = APP_DIR / "activation_request.json"
OCR_LANG = "por"

APP_NAME = "Separador de Listas e Provas"
APP_VERSION = "1.0"

AUTHOR_NAME = "Gilleade Gonçalves Silva"
AUTHOR_EMAIL = "gilleadesilva@gmail.com"
AUTHOR_LINKEDIN = "https://br.linkedin.com/in/gilleade-silva"
AUTHOR_GITHUB = "https://github.com/Gilleade"

AUTHOR_CREDIT_TEXT = (
    f"Desenvolvido por {AUTHOR_NAME} | "
    f"Contato: {AUTHOR_EMAIL}"
)

CODE_PATTERN = re.compile(r"\bPE-[0-9][A-Z&]+-[0-9]{5}\b", re.IGNORECASE)

LIST_HEADER_TERMS = [
    "REGISTRO DE TREINAMENTO",
    "CONSCIENTIZACAO",
    "CONSCIENTIZAÇÃO",
]

EXAM_HEADER_TERMS = [
    "AVALIACAO DA COMPREENSAO DO TREINAMENTO",
    "AVALIAÇÃO DA COMPREENSÃO DO TREINAMENTO",
]

LIST_BODY_TERMS = [
    "ASSUNTOS TRATADOS",
    "ASSINATURA",
    "FUNCAO",
    "FUNÇÃO",
    "PARTICIPANTE",
    "MATRICULA",
    "MATRÍCULA",
    "INSTRUTOR",
]

EXAM_BODY_TERMS = [
    "MARQUE VERDADEIRO OU FALSO",
    "VERDADEIRO",
    "FALSO",
    "NOTA",
    "QUESTAO",
    "QUESTÃO",
    "QUESTOES",
    "QUESTÕES",
    "CORRETA",
    "ASSINALE",
]

EMPLOYEE_LABEL_PATTERNS = [
    r"\bNOME\s*[:\-]?\s*([A-ZÀ-Ú][A-ZÀ-Ú\s]{4,120})",
    r"\bEMPREGADO\s*[:\-]?\s*([A-ZÀ-Ú][A-ZÀ-Ú\s]{4,120})",
    r"\bPARTICIPANTE\s*[:\-]?\s*([A-ZÀ-Ú][A-ZÀ-Ú\s]{4,120})",
    r"\bCOLABORADOR\s*[:\-]?\s*([A-ZÀ-Ú][A-ZÀ-Ú\s]{4,120})",
]

RAW_STOP_TOKENS = {
    "TREINAMENTO", "PROCEDIMENTO", "PETROBRAS", "AVALIACAO", "PROVA",
    "MANUAL", "SEGURANCA", "PRESENCA", "ANEXO", "FUNCAO", "FUNÇÃO",
    "ASSINATURA", "DATA", "NOTA", "EMPRESA", "MATRICULA", "MATRÍCULA",
    "INSTRUTOR", "CARGO", "NOME", "RG", "CPF", "SETOR", "UNIDADE"
}
STOPWORDS_NAME = {
    ''.join(c for c in unicodedata.normalize('NFD', x.upper()) if unicodedata.category(c) != 'Mn')
    for x in RAW_STOP_TOKENS
}


# ============================================================
# Estruturas de dados
# ============================================================

@dataclass
class PreAnalysisResult:
    file_name: str
    file_path: str
    total_pages: int
    is_probably_scanned: bool


@dataclass
class PageRecord:
    page_number: int
    raw_text: str
    norm_text: str
    page_type: str
    code: Optional[str] = None
    annex: Optional[str] = None
    logical_id: Optional[str] = None
    employee_name: Optional[str] = None
    confidence: float = 0.0
    notes: List[str] = field(default_factory=list)


@dataclass
class ListRecord:
    page_number: int
    code: Optional[str]
    annex: Optional[str]
    logical_id: Optional[str]
    confidence: float
    notes: List[str] = field(default_factory=list)


@dataclass
class ExamBlock:
    start_page: int
    pages: List[int]
    code: Optional[str]
    annex: Optional[str]
    logical_id: Optional[str]
    employee_name: Optional[str]
    confidence: float
    notes: List[str] = field(default_factory=list)


@dataclass
class PairRecord:
    logical_id: str
    list_record: Optional[ListRecord]
    exam_block: Optional[ExamBlock]
    status: str
    output_file: Optional[str] = None
    notes: List[str] = field(default_factory=list)
    

@dataclass
class ManualReviewItem:
    review_id: str
    pair: PairRecord
    list_record: Optional[ListRecord]
    suggested_exam_block: Optional[ExamBlock]
    candidate_exam_blocks: List[ExamBlock]
    typed_logical_id: str = ""
    review_status: str = "PENDENTE"
    confirmed_exam_start_page: Optional[int] = None
    reserved_exam_start_page: Optional[int] = None
    generated_output_file: Optional[str] = None
    notes: List[str] = field(default_factory=list)

@dataclass
class ProcessResult:
    input_file: str
    input_pdf_path: str
    collaborator_name: str
    total_pages: int
    page_records: List[PageRecord]
    lists_found: List[ListRecord]
    exams_found: List[ExamBlock]
    pairs: List[PairRecord]
    auto_export_pairs: List[PairRecord]
    manual_review_pairs: List[PairRecord]
    manual_review_items: List[ManualReviewItem]
    warnings: List[str]
    errors: List[str]
    output_dir: str
    report_txt_path: str
    report_csv_path: str


# ============================================================
# Utilitários gerais
# ============================================================

def ensure_base_folders() -> None:
    BRUTO_DIR.mkdir(parents=True, exist_ok=True)
    PROCESSADO_DIR.mkdir(parents=True, exist_ok=True)

def configure_tesseract() -> None:
    portable_exe = APP_DIR / "tesseract" / "tesseract.exe"
    system_exe = shutil_which("tesseract")

    if portable_exe.exists():
        pytesseract.pytesseract.tesseract_cmd = str(portable_exe)
        tessdata_dir = portable_exe.parent / "tessdata"
        if tessdata_dir.exists():
            os.environ.setdefault("TESSDATA_PREFIX", str(tessdata_dir))
        return

    if system_exe:
        pytesseract.pytesseract.tesseract_cmd = system_exe
        return

    raise RuntimeError(
        "Tesseract não encontrado. Coloque a pasta 'tesseract' ao lado do programa "
        "ou deixe o Tesseract disponível no PATH da máquina que vai gerar o pacote."
    )



def shutil_which(cmd: str) -> Optional[str]:
    paths = os.environ.get("PATH", "").split(os.pathsep)
    exts = [""]
    if os.name == "nt":
        exts.extend(os.environ.get("PATHEXT", ".EXE;.BAT;.CMD").split(";"))
    for path in paths:
        if not path:
            continue
        base = Path(path)
        for ext in exts:
            candidate = base / f"{cmd}{ext}"
            if candidate.exists():
                return str(candidate)
    return None



def normalize_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()



def remove_accents(text: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", text)
        if unicodedata.category(c) != "Mn"
    )



def sanitize_filename(text: str) -> str:
    text = remove_accents(text)
    text = re.sub(r'[\\/*?:"<>|]+', "", text)
    text = re.sub(r"\s+", "", text)
    return text.strip()

def build_author_report_lines() -> List[str]:
    return [
        f"Ferramenta desenvolvida por: {AUTHOR_NAME}",
        f"Contato: {AUTHOR_EMAIL}",
    ]


def build_about_text() -> str:
    return (
        f"{APP_NAME}\n"
        f"Versão: {APP_VERSION}\n\n"
        f"Desenvolvido por: {AUTHOR_NAME}\n"
        f"E-mail: {AUTHOR_EMAIL}\n"
        f"LinkedIn: {AUTHOR_LINKEDIN}\n"
        f"GitHub: {AUTHOR_GITHUB}"
    )

def contains_any(text: str, terms: List[str]) -> bool:
    return any(term in text for term in terms)



def is_success_status(status: str) -> bool:
    return str(status).startswith("OK")


def is_auto_export_status(status: str) -> bool:
    return str(status) in {"OK", "OK_BASE_CODE"}


def split_pairs_for_workflow(pairs: List[PairRecord]) -> Tuple[List[PairRecord], List[PairRecord]]:
    auto_export_pairs: List[PairRecord] = []
    manual_review_pairs: List[PairRecord] = []

    for pair in pairs:
        if is_auto_export_status(pair.status):
            auto_export_pairs.append(pair)
        else:
            manual_review_pairs.append(pair)

    return auto_export_pairs, manual_review_pairs


def build_manual_review_items(manual_review_pairs: List[PairRecord]) -> List[ManualReviewItem]:
    pending_exam_blocks: List[ExamBlock] = []
    seen_exam_pages = set()

    for pair in manual_review_pairs:
        if pair.exam_block and pair.exam_block.start_page not in seen_exam_pages:
            pending_exam_blocks.append(pair.exam_block)
            seen_exam_pages.add(pair.exam_block.start_page)

    review_items: List[ManualReviewItem] = []

    for idx, pair in enumerate(manual_review_pairs, start=1):
        suggested_exam = pair.exam_block

        candidate_exam_blocks = pending_exam_blocks.copy()

        if pair.list_record:
            candidate_exam_blocks.sort(
                key=lambda exam: (
                    0 if suggested_exam and exam.start_page == suggested_exam.start_page else 1,
                    abs(exam.start_page - pair.list_record.page_number),
                    exam.start_page,
                )
            )
        else:
            candidate_exam_blocks.sort(
                key=lambda exam: (
                    0 if suggested_exam and exam.start_page == suggested_exam.start_page else 1,
                    exam.start_page,
                )
            )

        review_id = f"REV_{idx:04d}"

        review_items.append(ManualReviewItem(
            review_id=review_id,
            pair=pair,
            list_record=pair.list_record,
            suggested_exam_block=suggested_exam,
            candidate_exam_blocks=candidate_exam_blocks,
            typed_logical_id=pair.logical_id or "",
            reserved_exam_start_page=None,
            notes=pair.notes.copy(),
        ))

    return review_items


def get_review_reserved_exam_pages(items: List[ManualReviewItem], exclude_review_id: Optional[str] = None) -> set:
    reserved = set()

    for item in items:
        if exclude_review_id and item.review_id == exclude_review_id:
            continue
        if item.reserved_exam_start_page is not None:
            reserved.add(item.reserved_exam_start_page)

    return reserved


def rebuild_candidate_exam_blocks(items: List[ManualReviewItem]) -> None:
    all_exam_map: Dict[int, ExamBlock] = {}

    for item in items:
        if item.suggested_exam_block:
            all_exam_map[item.suggested_exam_block.start_page] = item.suggested_exam_block
        for exam in item.candidate_exam_blocks:
            all_exam_map[exam.start_page] = exam
        if item.pair.exam_block:
            all_exam_map[item.pair.exam_block.start_page] = item.pair.exam_block

    all_exam_blocks = list(all_exam_map.values())

    for item in items:
        reserved_by_others = get_review_reserved_exam_pages(items, exclude_review_id=item.review_id)
        own_reserved_page = item.reserved_exam_start_page

        filtered = []
        for exam in all_exam_blocks:
            if exam.start_page == own_reserved_page:
                filtered.append(exam)
                continue
            if exam.start_page in reserved_by_others:
                continue
            filtered.append(exam)

        suggested_exam = item.suggested_exam_block

        if item.list_record:
            filtered.sort(
                key=lambda exam: (
                    0 if suggested_exam and exam.start_page == suggested_exam.start_page else 1,
                    0 if own_reserved_page is not None and exam.start_page == own_reserved_page else 1,
                    abs(exam.start_page - item.list_record.page_number),
                    exam.start_page,
                )
            )
        else:
            filtered.sort(
                key=lambda exam: (
                    0 if suggested_exam and exam.start_page == suggested_exam.start_page else 1,
                    0 if own_reserved_page is not None and exam.start_page == own_reserved_page else 1,
                    exam.start_page,
                )
            )

        item.candidate_exam_blocks = filtered


def reserve_exam_for_review_item(items: List[ManualReviewItem], target_item: ManualReviewItem, exam: ExamBlock) -> None:
    target_item.reserved_exam_start_page = exam.start_page
    target_item.confirmed_exam_start_page = exam.start_page
    target_item.pair.exam_block = exam
    rebuild_candidate_exam_blocks(items)


def release_exam_for_review_item(items: List[ManualReviewItem], target_item: ManualReviewItem, clear_pair_exam: bool = True) -> None:
    target_item.reserved_exam_start_page = None
    target_item.confirmed_exam_start_page = None
    if clear_pair_exam:
        target_item.pair.exam_block = None
    rebuild_candidate_exam_blocks(items)


# ============================================================
# PDF / OCR
# ============================================================

def open_pdf(file_path: str) -> fitz.Document:
    if not file_path or not os.path.exists(file_path):
        raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")
    return fitz.open(file_path)



def is_probably_scanned_pdf(doc: fitz.Document, sample_pages: int = 3) -> bool:
    total = min(sample_pages, len(doc))
    if total == 0:
        return True

    extracted_text_lengths = []
    for i in range(total):
        page = doc.load_page(i)
        text = page.get_text("text").strip()
        extracted_text_lengths.append(len(text))

    avg_len = sum(extracted_text_lengths) / len(extracted_text_lengths)
    return avg_len < 30



def pre_analyze_pdf(file_path: str) -> PreAnalysisResult:
    doc = open_pdf(file_path)
    try:
        return PreAnalysisResult(
            file_name=os.path.basename(file_path),
            file_path=file_path,
            total_pages=len(doc),
            is_probably_scanned=is_probably_scanned_pdf(doc),
        )
    finally:
        doc.close()



def render_page_to_pil(doc: fitz.Document, page_number: int, zoom: float = 2.0) -> Image.Image:
    page = doc.load_page(page_number)
    matrix = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=matrix, alpha=False)
    return Image.open(io.BytesIO(pix.tobytes("png")))



def preprocess_image_for_ocr(image: Image.Image) -> Image.Image:
    gray = ImageOps.grayscale(image)
    gray = ImageEnhance.Contrast(gray).enhance(1.8)
    gray = gray.filter(ImageFilter.SHARPEN)
    return gray



def normalize_for_matching(text: str) -> str:
    if not text:
        return ""
    text = remove_accents(text.upper())
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", "\n", text)
    return text.strip()



def extract_text_bundle_from_image(image: Image.Image) -> Tuple[str, str]:
    raw_text = pytesseract.image_to_string(image, lang=OCR_LANG) or ""
    norm_text = normalize_for_matching(raw_text)
    return raw_text, norm_text



def crop_vertical_region(image: Image.Image, start_ratio: float, end_ratio: float) -> Image.Image:
    width, height = image.size
    top = int(height * start_ratio)
    bottom = int(height * end_ratio)
    return image.crop((0, top, width, bottom))



def extract_page_region_texts(image: Image.Image) -> Dict[str, str]:
    full_processed = preprocess_image_for_ocr(image)
    full_raw, full_norm = extract_text_bundle_from_image(full_processed)

    header_img = crop_vertical_region(image, 0.0, 0.30)
    body_img = crop_vertical_region(image, 0.30, 0.92)

    header_processed = preprocess_image_for_ocr(header_img)
    body_processed = preprocess_image_for_ocr(body_img)

    header_raw, header_norm = extract_text_bundle_from_image(header_processed)
    body_raw, body_norm = extract_text_bundle_from_image(body_processed)

    return {
        "full_raw": full_raw,
        "full_norm": full_norm,
        "header_raw": header_raw,
        "header_norm": header_norm,
        "body_raw": body_raw,
        "body_norm": body_norm,
    }


# ============================================================
# Extrações e heurísticas
# ============================================================

def extract_code(text: str) -> Optional[str]:
    if not text:
        return None
    match = CODE_PATTERN.search(text.upper())
    return match.group(0).upper() if match else None



def extract_annex(text: str) -> Optional[str]:
    if not text:
        return None

    text_up = remove_accents(text.upper())
    text_up = normalize_spaces(text_up)

    patterns = [
        r"\bANEXO\s+([A-Z])\b",
        r"\(\s*ANEXO\s+([A-Z])\s*\)",
        r"\-\s*ANEXO\s+([A-Z])\b",
    ]

    for pattern in patterns:
        match = re.search(pattern, text_up, flags=re.IGNORECASE)
        if match:
            return f"ANEXO {match.group(1).upper()}"

    return None



def build_logical_identifier(code: Optional[str], annex: Optional[str]) -> Optional[str]:
    if not code:
        return None
    if annex:
        annex_norm = annex.replace(" ", "_")
        return f"{code}__{annex_norm}"
    return code



def extract_code_flexible(norm_text: str) -> Optional[str]:
    if not norm_text:
        return None

    direct = extract_code(norm_text)
    if direct:
        return direct

    repaired_text = norm_text.replace(" ", "")
    repaired_patterns = [
        r"PE-[0-9][A-Z&]+-[0-9]{5}",
        r"PE-[0-9][A-Z&]+[0-9]{5}",
    ]

    for pattern in repaired_patterns:
        match = re.search(pattern, repaired_text, flags=re.IGNORECASE)
        if match:
            value = match.group(0).upper()
            if value.count("-") == 1:
                prefix = re.match(r"(PE-[0-9][A-Z&]+)([0-9]{5})", value)
                if prefix:
                    value = f"{prefix.group(1)}-{prefix.group(2)}"
            return value
    return None


def extract_code_from_regions(
    header_norm_text: str,
    full_norm_text: str,
    body_norm_text: str,
) -> Optional[str]:
    candidates = [
        header_norm_text or "",
        full_norm_text or "",
        body_norm_text or "",
    ]

    for text in candidates:
        code = extract_code_flexible(text)
        if code:
            return code

    return None


def has_list_table_signature(body_norm: str) -> bool:
    markers = 0
    for term in ["NOME", "FUNCAO", "ASSINATURA"]:
        if term in body_norm:
            markers += 1
    return markers >= 2



def count_question_markers(text: str) -> int:
    if not text:
        return 0

    patterns = [
        r"(?m)(?:^|\s)([0-9]{1,2})\s*[\.\)\-–—]",
        r"(?:^|\s)([A-D])[\.\)]",
        r"\bASSINALE\b",
        r"\bMARQUE\b",
        r"\bCORRETA\b",
        r"\bQUESTAO\b",
        r"\bQUESTÃO\b",
        r"\bQUESTOES\b",
        r"\bQUESTÕES\b",
    ]

    total = 0
    for pattern in patterns:
        total += len(re.findall(pattern, text, flags=re.IGNORECASE | re.MULTILINE))
    return total



def has_exam_question_structure(body_norm: str) -> bool:
    qmarks = count_question_markers(body_norm)
    if qmarks >= 2:
        return True
    if "VERDADEIRO" in body_norm and "FALSO" in body_norm:
        return True
    if contains_any(body_norm, EXAM_BODY_TERMS):
        return True
    return False



def normalize_token_for_name(token: str) -> str:
    return remove_accents(token.upper()).strip()



def clean_candidate_name(name: str) -> Optional[str]:
    if not name:
        return None

    name = normalize_spaces(name.replace("\n", " "))
    name = re.sub(r"[^A-Za-zÀ-ú\s]", " ", name)
    name = normalize_spaces(name)

    if len(name) < 5:
        return None

    raw_parts = [p for p in name.split() if p.strip()]
    cleaned_parts = []

    for part in raw_parts:
        token_norm = normalize_token_for_name(part)
        if token_norm in STOPWORDS_NAME:
            break
        if any(ch.isdigit() for ch in part):
            break
        if len(part) == 1:
            continue
        cleaned_parts.append(part)
        if len(cleaned_parts) >= 6:
            break

    if len(cleaned_parts) < 2:
        return None

    candidate = " ".join(cleaned_parts).title().strip()
    return candidate if len(candidate) >= 5 else None



def extract_employee_name(raw_text: str, norm_text: str) -> Optional[str]:
    for pattern in EMPLOYEE_LABEL_PATTERNS:
        match = re.search(pattern, raw_text, flags=re.IGNORECASE)
        if match:
            candidate = clean_candidate_name(match.group(1))
            if candidate:
                return candidate

    lines = [normalize_spaces(line) for line in raw_text.splitlines() if normalize_spaces(line)]
    for line in lines:
        line_up = remove_accents(line.upper())
        if any(label in line_up for label in ["NOME", "EMPREGADO", "PARTICIPANTE", "COLABORADOR"]):
            maybe = re.split(r"[:\-]", line, maxsplit=1)
            if len(maybe) == 2:
                candidate = clean_candidate_name(maybe[1])
                if candidate:
                    return candidate
    return None



def parse_code_parts(code: Optional[str]) -> Optional[Dict[str, str]]:
    if not code:
        return None
    m = re.match(r"^PE-([0-9])([A-Z&]+)-([0-9]{5})$", code.upper())
    if not m:
        return None
    return {
        "zone_digit": m.group(1),
        "alpha": m.group(2),
        "serial": m.group(3),
    }



def parse_code_parts_extended(code: Optional[str]) -> Optional[Dict[str, str]]:
    if not code:
        return None
    code = code.upper().strip()
    m = re.match(r"^PE-([0-9])([A-Z&]+)-([0-9]{5})$", code)
    if not m:
        return None
    alpha = m.group(2)
    return {
        "zone_digit": m.group(1),
        "alpha": alpha,
        "alpha_compact": compress_repeated_chars(alpha),
        "serial": m.group(3),
    }



def compress_repeated_chars(text: str) -> str:
    if not text:
        return text
    out: List[str] = []
    for ch in text:
        if not out or out[-1] != ch:
            out.append(ch)
    return "".join(out)



def alpha_similarity(a: str, b: str) -> float:
    return SequenceMatcher(None, a, b).ratio()



def codes_are_soft_compatible(code_a: Optional[str], code_b: Optional[str]) -> bool:
    if not code_a or not code_b:
        return False
    if code_a == code_b:
        return True
    a = parse_code_parts_extended(code_a)
    b = parse_code_parts_extended(code_b)
    if not a or not b:
        return False
    if a["serial"] == b["serial"]:
        sim_raw = alpha_similarity(a["alpha"], b["alpha"])
        sim_compact = alpha_similarity(a["alpha_compact"], b["alpha_compact"])
        if sim_raw >= 0.80 or sim_compact >= 0.80:
            return True
        if a["alpha_compact"] == b["alpha_compact"]:
            return True
    return False



def annexes_soft_compatible(annex_a: Optional[str], annex_b: Optional[str]) -> bool:
    if (annex_a or "") == (annex_b or ""):
        return True
    if not annex_a or not annex_b:
        return True
    return False



def is_strong_list_layout(norm_text: str) -> bool:
    markers = 0
    strong_terms = [
        "LISTA DE PRESENCA", "PRESENCA", "ASSINATURA", "FUNCAO",
        "PARTICIPANTE", "MATRICULA", "EMPRESA", "INSTRUTOR",
    ]
    for term in strong_terms:
        if term in norm_text:
            markers += 1
    if "LISTA" in norm_text and "PRESENCA" in norm_text:
        markers += 2
    col_terms = ["NOME", "FUNCAO", "ASSINATURA"]
    if sum(1 for term in col_terms if term in norm_text) >= 2:
        markers += 2
    return markers >= 4



def is_exam_like(norm_text: str) -> bool:
    score = 0
    for term in ["AVALIACAO", "PROVA", "NOTA", "ANEXO", "QUESTAO", "QUESTOES", "CORRETA"]:
        if term in norm_text:
            score += 1
    if count_question_markers(norm_text) >= 2:
        score += 2
    return score >= 3



def score_list_page(norm_text: str) -> float:
    score = 0.0
    positive_terms = [
        "LISTA DE PRESENCA", "PRESENCA", "ASSINATURA", "PARTICIPANTE",
        "FUNCAO", "MATRICULA", "EMPRESA", "INSTRUTOR",
    ]
    for term in positive_terms:
        if term in norm_text:
            score += 1.3
    if "LISTA" in norm_text and "PRESENCA" in norm_text:
        score += 2.5
    if is_strong_list_layout(norm_text):
        score += 3.0
    if extract_code_flexible(norm_text):
        score += 0.8
    if count_question_markers(norm_text) >= 2:
        score -= 1.5
    return score



def score_exam_start_page(norm_text: str) -> float:
    score = 0.0
    positive_terms = ["AVALIACAO", "PROVA", "TREINAMENTO", "NOME", "NOTA", "DATA", "ANEXO"]
    for term in positive_terms:
        if term in norm_text:
            score += 1.0
    if extract_code_flexible(norm_text):
        score += 1.5
    qm = count_question_markers(norm_text)
    if qm >= 2:
        score += 1.8
    elif qm == 1:
        score += 0.8
    if is_strong_list_layout(norm_text):
        score -= 3.0
    return score



def score_exam_continuation_page(norm_text: str) -> float:
    score = 0.0
    qm = count_question_markers(norm_text)
    if qm >= 3:
        score += 2.5
    elif qm >= 1:
        score += 1.3
    for term in ["ASSINALE", "CORRETA", "QUESTAO", "QUESTOES", "MARQUE"]:
        if term in norm_text:
            score += 0.8
    if extract_code_flexible(norm_text):
        score -= 0.5
    if is_strong_list_layout(norm_text):
        score -= 3.0
    return score



def classify_page(
    raw_text: str,
    norm_text: str,
    previous_type: Optional[str] = None,
    header_norm_text: str = "",
    body_norm_text: str = "",
) -> Tuple[str, float, List[str]]:
    notes: List[str] = []

    if contains_any(header_norm_text, EXAM_HEADER_TERMS):
        notes.append("cabecalho_prova_detectado")
        return "PROVA_INICIO", 0.99, notes

    if contains_any(header_norm_text, LIST_HEADER_TERMS):
        notes.append("cabecalho_lista_detectado")
        return "LISTA", 0.99, notes

    if "ASSUNTOS TRATADOS" in body_norm_text:
        notes.append("campo_assuntos_tratados_detectado")
        return "LISTA", 0.97, notes

    if has_list_table_signature(body_norm_text) and contains_any(body_norm_text, LIST_BODY_TERMS):
        notes.append("estrutura_lista_detectada")
        return "LISTA", 0.96, notes

    if has_exam_question_structure(body_norm_text):
        if previous_type in {"PROVA_INICIO", "PROVA_CONTINUACAO"}:
            notes.append("estrutura_questoes_detectada_continuacao")
            return "PROVA_CONTINUACAO", 0.94, notes
        notes.append("estrutura_questoes_detectada_inicio")
        return "PROVA_INICIO", 0.94, notes

    list_score = score_list_page(norm_text)
    exam_start_score = score_exam_start_page(norm_text)
    exam_cont_score = score_exam_continuation_page(norm_text)

    if is_strong_list_layout(norm_text) and list_score >= 4.0:
        notes.append(f"fallback_score_lista={list_score:.1f}")
        return "LISTA", min(0.92, 0.60 + list_score / 10), notes

    if list_score >= 3.4 and list_score > exam_start_score:
        notes.append(f"fallback_score_lista={list_score:.1f}")
        return "LISTA", min(0.90, 0.55 + list_score / 10), notes

    if exam_start_score >= 3.2 and exam_start_score >= exam_cont_score:
        notes.append(f"fallback_score_prova_inicio={exam_start_score:.1f}")
        return "PROVA_INICIO", min(0.90, 0.55 + exam_start_score / 10), notes

    if previous_type in {"PROVA_INICIO", "PROVA_CONTINUACAO"}:
        if not is_strong_list_layout(norm_text) and exam_cont_score >= 1.0:
            notes.append(f"fallback_score_prova_cont={exam_cont_score:.1f}")
            return "PROVA_CONTINUACAO", min(0.88, 0.45 + exam_cont_score / 10), notes

    notes.append(f"indefinida(lista={list_score:.1f}, inicio={exam_start_score:.1f}, cont={exam_cont_score:.1f})")
    return "INDEFINIDA", 0.30, notes


# ============================================================
# Pipeline principal
# ============================================================

def analyze_pdf_pages(file_path: str, progress_callback: Optional[Callable[[str], None]] = None) -> List[PageRecord]:
    doc = open_pdf(file_path)
    page_records: List[PageRecord] = []
    previous_type = None

    try:
        for page_idx in range(len(doc)):
            if progress_callback:
                progress_callback(f"Lendo página {page_idx + 1} de {len(doc)}...")

            pil_image = render_page_to_pil(doc, page_idx)
            region_texts = extract_page_region_texts(pil_image)

            raw_text = region_texts["full_raw"]
            norm_text = region_texts["full_norm"]
            header_norm_text = region_texts["header_norm"]
            body_norm_text = region_texts["body_norm"]
            code = extract_code_from_regions(
                header_norm_text=header_norm_text,
                full_norm_text=norm_text,
                body_norm_text=body_norm_text,
            )

            page_type, confidence, notes = classify_page(
                raw_text=raw_text,
                norm_text=norm_text,
                previous_type=previous_type,
                header_norm_text=header_norm_text,
                body_norm_text=body_norm_text,
            )

            annex = extract_annex(header_norm_text)
            if not annex:
                annex = extract_annex(norm_text)

            if not annex and page_type == "LISTA":
                annex = extract_annex(body_norm_text)

            logical_id = build_logical_identifier(code, annex)
            employee_name = None
            if page_type in {"PROVA_INICIO", "PROVA_CONTINUACAO"}:
                employee_name = extract_employee_name(raw_text, norm_text)

            notes = notes + [
                f"header={header_norm_text[:120]}",
                f"body={body_norm_text[:120]}",
            ]

            page_records.append(PageRecord(
                page_number=page_idx + 1,
                raw_text=raw_text,
                norm_text=norm_text,
                page_type=page_type,
                code=code,
                annex=annex,
                logical_id=logical_id,
                employee_name=employee_name,
                confidence=confidence,
                notes=notes,
            ))
            previous_type = page_type
    finally:
        doc.close()

    return page_records



def collect_lists(page_records: List[PageRecord]) -> List[ListRecord]:
    results: List[ListRecord] = []
    for page in page_records:
        if page.page_type == "LISTA":
            results.append(ListRecord(
                page_number=page.page_number,
                code=page.code,
                annex=page.annex,
                logical_id=page.logical_id,
                confidence=page.confidence,
                notes=page.notes.copy(),
            ))
    return results



def most_common_non_null(values: List[Optional[str]]) -> Optional[str]:
    cleaned = [v for v in values if v]
    if not cleaned:
        return None
    return Counter(cleaned).most_common(1)[0][0]



def should_absorb_undefined_into_exam(page: PageRecord) -> bool:
    if page.page_type != "INDEFINIDA":
        return False
    if is_strong_list_layout(page.norm_text):
        return False
    if count_question_markers(page.norm_text) >= 1:
        return True
    if is_exam_like(page.norm_text):
        return True
    return False



def finalize_exam_block(page_records: List[PageRecord], current_block_pages: List[int], current_notes: List[str]) -> ExamBlock:
    block_records = [page_records[p - 1] for p in current_block_pages]
    code = most_common_non_null([r.code for r in block_records])
    annex = most_common_non_null([r.annex for r in block_records])
    logical_id = most_common_non_null([r.logical_id for r in block_records])

    employee_name = None
    for rec in block_records:
        if rec.page_type == "PROVA_INICIO" and rec.employee_name:
            employee_name = rec.employee_name
            break
    if not employee_name:
        employee_name = most_common_non_null([r.employee_name for r in block_records])

    return ExamBlock(
        start_page=current_block_pages[0],
        pages=current_block_pages.copy(),
        code=code,
        annex=annex,
        logical_id=logical_id,
        employee_name=employee_name,
        confidence=sum(r.confidence for r in block_records) / len(block_records),
        notes=current_notes.copy(),
    )



def group_exam_blocks(page_records: List[PageRecord]) -> List[ExamBlock]:
    exam_blocks: List[ExamBlock] = []
    current_block_pages: List[int] = []
    current_notes: List[str] = []

    for page in page_records:
        if page.page_type == "PROVA_INICIO":
            if current_block_pages:
                exam_blocks.append(finalize_exam_block(page_records, current_block_pages, current_notes))
                current_block_pages = []
                current_notes = []
            current_block_pages = [page.page_number]
            current_notes = page.notes.copy()

        elif page.page_type == "PROVA_CONTINUACAO":
            if current_block_pages:
                current_block_pages.append(page.page_number)
                current_notes.extend(page.notes)

        elif current_block_pages and should_absorb_undefined_into_exam(page):
            current_block_pages.append(page.page_number)
            current_notes.extend(page.notes + ["pagina_indefinida_absorvida_como_continuacao"])

        else:
            if current_block_pages:
                exam_blocks.append(finalize_exam_block(page_records, current_block_pages, current_notes))
                current_block_pages = []
                current_notes = []

    if current_block_pages:
        exam_blocks.append(finalize_exam_block(page_records, current_block_pages, current_notes))

    return exam_blocks



def split_logical_id(logical_id: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
    if not logical_id:
        return None, None
    if "__" in logical_id:
        base, annex = logical_id.split("__", 1)
        return base, annex
    return logical_id, None



def base_code_from_logical_id(logical_id: Optional[str]) -> Optional[str]:
    base, _ = split_logical_id(logical_id)
    return base



def annex_from_logical_id(logical_id: Optional[str]) -> Optional[str]:
    _, annex = split_logical_id(logical_id)
    return annex



def pages_are_close(list_page: int, exam_page: int, max_distance: int = 2) -> bool:
    return abs(list_page - exam_page) <= max_distance



def make_pair(
    logical_id: str,
    list_record: Optional[ListRecord],
    exam_block: Optional[ExamBlock],
    status: str,
    notes: Optional[List[str]] = None,
) -> PairRecord:
    return PairRecord(
        logical_id=logical_id,
        list_record=list_record,
        exam_block=exam_block,
        status=status,
        notes=notes or [],
    )



def pick_preferred_logical_id(list_item: Optional[ListRecord], exam_item: Optional[ExamBlock]) -> str:
    if exam_item and exam_item.logical_id and "__ANEXO_" in exam_item.logical_id:
        return exam_item.logical_id
    if list_item and list_item.logical_id:
        return list_item.logical_id
    if exam_item and exam_item.logical_id:
        return exam_item.logical_id
    if list_item and list_item.code:
        return list_item.code
    if exam_item and exam_item.code:
        return exam_item.code
    if list_item:
        return f"SEM_ID_LISTA_PAG_{list_item.page_number}"
    if exam_item:
        return f"SEM_ID_PROVA_PAG_{exam_item.start_page}"
    return "SEM_IDENTIFICADOR"



def pair_lists_and_exams(lists_found: List[ListRecord], exams_found: List[ExamBlock]) -> Tuple[List[PairRecord], List[str], List[str]]:
    warnings: List[str] = []
    errors: List[str] = []
    pairs: List[PairRecord] = []

    remaining_lists = lists_found.copy()
    remaining_exams = exams_found.copy()

    def remove_matched(list_item=None, exam_item=None):
        nonlocal remaining_lists, remaining_exams
        if list_item in remaining_lists:
            remaining_lists.remove(list_item)
        if exam_item in remaining_exams:
            remaining_exams.remove(exam_item)

    paired_now = True
    while paired_now:
        paired_now = False
        list_map = defaultdict(list)
        exam_map = defaultdict(list)
        for item in remaining_lists:
            key = item.logical_id or f"SEM_ID_LISTA_PAG_{item.page_number}"
            list_map[key].append(item)
        for item in remaining_exams:
            key = item.logical_id or f"SEM_ID_PROVA_PAG_{item.start_page}"
            exam_map[key].append(item)
        common_ids = sorted(set(list_map.keys()) & set(exam_map.keys()))
        for logical_id in common_ids:
            list_items = list_map[logical_id]
            exam_items = exam_map[logical_id]
            if len(list_items) == 1 and len(exam_items) == 1:
                list_item = list_items[0]
                exam_item = exam_items[0]
                pairs.append(make_pair(
                    logical_id=logical_id,
                    list_record=list_item,
                    exam_block=exam_item,
                    status="OK",
                    notes=["Pareamento exato por identificador lógico."],
                ))
                remove_matched(list_item, exam_item)
                paired_now = True
                break

    paired_now = True
    while paired_now:
        paired_now = False
        list_groups = defaultdict(list)
        exam_groups = defaultdict(list)
        for item in remaining_lists:
            base = base_code_from_logical_id(item.logical_id) or item.code
            if base:
                list_groups[base].append(item)
        for item in remaining_exams:
            base = base_code_from_logical_id(item.logical_id) or item.code
            if base:
                exam_groups[base].append(item)
        for base_code in sorted(set(list_groups.keys()) & set(exam_groups.keys())):
            list_items = list_groups[base_code]
            exam_items = exam_groups[base_code]
            normal_exams = []
            for ex in exam_items:
                ex_annex = annex_from_logical_id(ex.logical_id) or ex.annex
                if not ex_annex:
                    normal_exams.append(ex)
            if len(list_items) >= 2 and len(normal_exams) == 1:
                exam_item = normal_exams[0]
                list_item = min(list_items, key=lambda x: abs(x.page_number - exam_item.start_page))
                logical_id = pick_preferred_logical_id(list_item, exam_item)
                pairs.append(make_pair(
                    logical_id=logical_id,
                    list_record=list_item,
                    exam_block=exam_item,
                    status="OK_DESEMPATE_PROXIMIDADE",
                    notes=["Desempate por proximidade entre múltiplas listas do mesmo código-base."],
                ))
                warnings.append(
                    f"Desempate por proximidade: lista pág. {list_item.page_number} com prova pág. {exam_item.start_page} para {base_code}."
                )
                remove_matched(list_item, exam_item)
                paired_now = True
                break

    paired_now = True
    while paired_now:
        paired_now = False
        for list_item in remaining_lists[:]:
            list_base = base_code_from_logical_id(list_item.logical_id) or list_item.code
            list_annex = annex_from_logical_id(list_item.logical_id) or list_item.annex
            if not list_base:
                continue
            candidates = []
            for exam_item in remaining_exams:
                exam_base = base_code_from_logical_id(exam_item.logical_id) or exam_item.code
                exam_annex = annex_from_logical_id(exam_item.logical_id) or exam_item.annex
                if not exam_base:
                    continue
                if list_base == exam_base and annexes_soft_compatible(list_annex, exam_annex):
                    candidates.append(exam_item)
            if len(candidates) == 1:
                exam_item = candidates[0]
                logical_id = pick_preferred_logical_id(list_item, exam_item)
                pairs.append(make_pair(
                    logical_id=logical_id,
                    list_record=list_item,
                    exam_block=exam_item,
                    status="OK_BASE_CODE",
                    notes=["Pareamento por código-base, ignorando diferença de anexo."],
                ))
                remove_matched(list_item, exam_item)
                paired_now = True
                break

    paired_now = True
    while paired_now:
        paired_now = False
        for list_item in remaining_lists[:]:
            if not list_item.code:
                continue
            candidates = []
            for exam_item in remaining_exams:
                if not exam_item.code:
                    continue
                if annexes_soft_compatible(list_item.annex, exam_item.annex) and codes_are_soft_compatible(list_item.code, exam_item.code):
                    candidates.append(exam_item)
            if len(candidates) == 1:
                exam_item = candidates[0]
                logical_id = list_item.code or exam_item.code or pick_preferred_logical_id(list_item, exam_item)
                pairs.append(make_pair(
                    logical_id=logical_id,
                    list_record=list_item,
                    exam_block=exam_item,
                    status="OK_AJUSTE_OCR",
                    notes=[f"Pareamento por similaridade OCR: lista {list_item.code} x prova {exam_item.code}."],
                ))
                warnings.append(
                    f"Pareamento por similaridade OCR: lista pág. {list_item.page_number} ({list_item.code}) com prova pág. {exam_item.start_page} ({exam_item.code})."
                )
                remove_matched(list_item, exam_item)
                paired_now = True
                break

    paired_now = True
    while paired_now:
        paired_now = False
        for list_item in remaining_lists[:]:
            if list_item.code:
                continue

            candidates = []
            for exam_item in remaining_exams:
                if not exam_item.code:
                    continue
                if pages_are_close(list_item.page_number, exam_item.start_page, max_distance=2):
                    candidates.append(exam_item)

            if len(candidates) == 1:
                exam_item = candidates[0]
                logical_id = pick_preferred_logical_id(list_item, exam_item)
                pairs.append(make_pair(
                    logical_id=logical_id,
                    list_record=list_item,
                    exam_block=exam_item,
                    status="OK_PROXIMIDADE",
                    notes=["Pareamento por proximidade: lista sem código e prova próxima com código."],
                ))
                warnings.append(
                    f"Pareamento por proximidade: lista pág. {list_item.page_number} com prova pág. {exam_item.start_page}."
                )
                remove_matched(list_item, exam_item)
                paired_now = True
                break

    paired_now = True
    while paired_now:
        paired_now = False
        for list_item in remaining_lists[:]:
            if not list_item.code:
                continue

            candidates = []
            for exam_item in remaining_exams:
                if exam_item.code:
                    continue
                if pages_are_close(list_item.page_number, exam_item.start_page, max_distance=2):
                    candidates.append(exam_item)

            if len(candidates) == 1:
                exam_item = candidates[0]
                logical_id = list_item.code or pick_preferred_logical_id(list_item, exam_item)
                pairs.append(make_pair(
                    logical_id=logical_id,
                    list_record=list_item,
                    exam_block=exam_item,
                    status="REVISAO_PROXIMIDADE_CODIGO",
                    notes=["Pareamento sugerido por proximidade: prova sem código e lista com código próximo; revisão manual necessária."],
                ))
                warnings.append(
                    f"Revisão necessária (proximidade com código): lista pág. {list_item.page_number} ({list_item.code}) com prova pág. {exam_item.start_page}."
                )
                remove_matched(list_item, exam_item)
                paired_now = True
                break

    remaining_lists.sort(key=lambda x: x.page_number)
    remaining_exams.sort(key=lambda x: x.start_page)
    used_exam_indices = set()

    for list_item in remaining_lists[:]:
        best_idx = None
        best_score = None

        for idx, exam_item in enumerate(remaining_exams):
            if idx in used_exam_indices:
                continue

            distance = abs(list_item.page_number - exam_item.start_page)
            penalty = 0

            if list_item.code and exam_item.code and codes_are_soft_compatible(list_item.code, exam_item.code):
                penalty -= 8
            if (not list_item.code) or (not exam_item.code):
                penalty -= 2
            if exam_item.start_page >= list_item.page_number:
                penalty -= 1

            score = distance + penalty

            if distance > 20:
                continue

            if best_score is None or score < best_score:
                best_score = score
                best_idx = idx

        if best_idx is not None:
            exam_item = remaining_exams[best_idx]
            used_exam_indices.add(best_idx)

            logical_id = pick_preferred_logical_id(list_item, exam_item)
            pairs.append(make_pair(
                logical_id=logical_id,
                list_record=list_item,
                exam_block=exam_item,
                status="REVISAO_FALLBACK_ORDEM",
                notes=["Pareamento sugerido por ordem e distância entre páginas; revisão manual necessária."],
            ))
            warnings.append(
                f"Revisão necessária (fallback por ordem): lista pág. {list_item.page_number} com prova pág. {exam_item.start_page}."
            )

    paired_list_pages = {p.list_record.page_number for p in pairs if p.list_record}
    paired_exam_pages = {p.exam_block.start_page for p in pairs if p.exam_block}

    final_remaining_lists = [item for item in remaining_lists if item.page_number not in paired_list_pages]
    final_remaining_exams = [item for item in remaining_exams if item.start_page not in paired_exam_pages]

    for list_item in final_remaining_lists:
        logical_id = list_item.logical_id or f"SEM_ID_LISTA_PAG_{list_item.page_number}"
        pairs.append(make_pair(
            logical_id=logical_id,
            list_record=list_item,
            exam_block=None,
            status="LISTA_SEM_PROVA",
            notes=["Lista encontrada sem prova correspondente."],
        ))
        warnings.append(f"Lista sem prova para identificador: {logical_id}")

    for exam_item in final_remaining_exams:
        logical_id = exam_item.logical_id or f"SEM_ID_PROVA_PAG_{exam_item.start_page}"
        pairs.append(make_pair(
            logical_id=logical_id,
            list_record=None,
            exam_block=exam_item,
            status="PROVA_SEM_LISTA",
            notes=["Prova encontrada sem lista correspondente."],
        ))
        warnings.append(f"Prova sem lista para identificador: {logical_id}")

    return pairs, warnings, errors


def extract_name_from_input_filename(file_path: str) -> Optional[str]:
    if not file_path:
        return None

    stem = Path(file_path).stem
    stem = remove_accents(stem.upper())
    stem = re.sub(r"[_\-]+", " ", stem)
    stem = re.sub(r"\b(KIT|PDF|LISTA|PROVA|RELATORIO|RESUMO)\b", " ", stem)
    stem = normalize_spaces(stem)

    if not stem:
        return None

    candidate = clean_candidate_name(stem.title())
    return candidate


def choose_global_collaborator_name(
    file_path: str,
    exams_found: List[ExamBlock],
    page_records: List[PageRecord]
) -> Optional[str]:
    filename_name = extract_name_from_input_filename(file_path)
    if filename_name:
        return filename_name

    exam_names = [x.employee_name for x in exams_found if x.employee_name]
    if exam_names:
        return Counter(exam_names).most_common(1)[0][0]

    page_names = [
        x.employee_name for x in page_records
        if x.employee_name and x.page_type in {"PROVA_INICIO", "PROVA_CONTINUACAO"}
    ]
    if page_names:
        return Counter(page_names).most_common(1)[0][0]

    return None


def split_first_last_name(full_name: str) -> Tuple[str, str]:
    if not full_name:
        return ("Colaborador", "NaoIdentificado")
    parts = [p for p in normalize_spaces(full_name).split() if p]
    if len(parts) == 1:
        return parts[0], parts[0]
    return parts[0], parts[-1]



def build_collaborator_folder_name(full_name: Optional[str]) -> str:
    if not full_name:
        return "ColaboradorNaoIdentificado"
    first_name, last_name = split_first_last_name(full_name)
    return sanitize_filename(f"{first_name}{last_name}") or "ColaboradorNaoIdentificado"



def build_output_filename(full_name: Optional[str], logical_id: Optional[str]) -> str:
    base_name = build_collaborator_folder_name(full_name)
    logical_part = (logical_id or "SEM_IDENTIFICADOR").replace("__", "_")
    logical_part = sanitize_filename(logical_part)
    return f"{base_name}_{logical_part}.pdf"


def normalize_manual_logical_id(text: str) -> str:
    value = remove_accents((text or "").upper())
    value = normalize_spaces(value)
    value = re.sub(r"\s+ANEXO\s+([A-Z])\b", r"__ANEXO_\1", value)
    value = value.replace("/", "-")
    value = re.sub(r"\s+", "", value)
    return value.strip()


def resolve_manual_review_logical_id(item: ManualReviewItem, exam_block: Optional[ExamBlock]) -> str:
    typed = normalize_manual_logical_id(item.typed_logical_id)
    if typed:
        return typed
    return pick_preferred_logical_id(item.list_record, exam_block)


def ensure_output_structure(collaborator_name: Optional[str]) -> Tuple[str, str, str]:
    collaborator_folder = build_collaborator_folder_name(collaborator_name)
    run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
    collaborator_root = PROCESSADO_DIR / collaborator_folder
    run_output_dir = collaborator_root / run_id
    pdfs_dir = run_output_dir / "pdfs"
    collaborator_root.mkdir(parents=True, exist_ok=True)
    run_output_dir.mkdir(parents=True, exist_ok=True)
    pdfs_dir.mkdir(parents=True, exist_ok=True)
    return str(run_output_dir), str(pdfs_dir), run_id

def export_pair_pdf(input_pdf_path: str, pair: PairRecord, output_pdf_path: str) -> None:
    source_doc = fitz.open(input_pdf_path)
    out_doc = fitz.open()
    try:
        if pair.list_record:
            idx = pair.list_record.page_number - 1
            out_doc.insert_pdf(source_doc, from_page=idx, to_page=idx)
        if pair.exam_block:
            for page_number in pair.exam_block.pages:
                idx = page_number - 1
                out_doc.insert_pdf(source_doc, from_page=idx, to_page=idx)
        out_doc.save(output_pdf_path)
    finally:
        out_doc.close()
        source_doc.close()
        
PREVIEW_IMAGE_CACHE: Dict[Tuple[str, int, float], Image.Image] = {}
PREVIEW_PAGE_SIZE_CACHE: Dict[Tuple[str, int], Tuple[float, float]] = {}

MAX_PREVIEW_IMAGE_CACHE = 48
MAX_PREVIEW_PAGE_SIZE_CACHE = 256


def _cache_put_image(key: Tuple[str, int, float], image: Image.Image) -> None:
    PREVIEW_IMAGE_CACHE[key] = image
    if len(PREVIEW_IMAGE_CACHE) > MAX_PREVIEW_IMAGE_CACHE:
        oldest_key = next(iter(PREVIEW_IMAGE_CACHE))
        PREVIEW_IMAGE_CACHE.pop(oldest_key, None)


def _cache_put_page_size(key: Tuple[str, int], size: Tuple[float, float]) -> None:
    PREVIEW_PAGE_SIZE_CACHE[key] = size
    if len(PREVIEW_PAGE_SIZE_CACHE) > MAX_PREVIEW_PAGE_SIZE_CACHE:
        oldest_key = next(iter(PREVIEW_PAGE_SIZE_CACHE))
        PREVIEW_PAGE_SIZE_CACHE.pop(oldest_key, None)


def clear_preview_caches() -> None:
    PREVIEW_IMAGE_CACHE.clear()
    PREVIEW_PAGE_SIZE_CACHE.clear()

def render_pdf_page_preview(
    input_pdf_path: str,
    page_number: int,
    max_size: Optional[Tuple[int, int]] = None,
    zoom: float = 1.35,
) -> Image.Image:
    if not input_pdf_path:
        raise ValueError("Caminho do PDF de origem não informado.")

    if page_number < 1:
        raise ValueError("Número de página inválido para pré-visualização.")

    normalized_path = str(Path(input_pdf_path).resolve())
    cache_key = (normalized_path, page_number, round(float(zoom), 2))

    cached = PREVIEW_IMAGE_CACHE.get(cache_key)
    if cached is not None:
        image = cached.copy()
    else:
        doc = fitz.open(input_pdf_path)
        try:
            page = doc.load_page(page_number - 1)
            matrix = fitz.Matrix(zoom, zoom)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            _cache_put_image(cache_key, image.copy())
        finally:
            doc.close()

    if max_size:
        resampling_module = getattr(Image, "Resampling", Image)
        image.thumbnail(max_size, resampling_module.LANCZOS)

    return image
        
def get_pdf_page_size(input_pdf_path: str, page_number: int) -> Tuple[float, float]:
    if not input_pdf_path:
        raise ValueError("Caminho do PDF de origem não informado.")

    if page_number < 1:
        raise ValueError("Número de página inválido.")

    normalized_path = str(Path(input_pdf_path).resolve())
    cache_key = (normalized_path, page_number)

    cached = PREVIEW_PAGE_SIZE_CACHE.get(cache_key)
    if cached is not None:
        return cached

    doc = fitz.open(input_pdf_path)
    try:
        page = doc.load_page(page_number - 1)
        rect = page.rect
        size = (float(rect.width), float(rect.height))
        _cache_put_page_size(cache_key, size)
        return size
    finally:
        doc.close()

def ensure_unique_output_path(output_path: str) -> str:
    path = Path(output_path)
    if not path.exists():
        return str(path)
    counter = 2
    while True:
        candidate = path.with_name(f"{path.stem}_{counter}{path.suffix}")
        if not candidate.exists():
            return str(candidate)
        counter += 1


def export_ok_pairs(input_pdf_path: str, pairs: List[PairRecord], pdfs_dir: str, fallback_name: Optional[str]) -> List[PairRecord]:
    updated_pairs: List[PairRecord] = []
    for pair in pairs:
        if not is_auto_export_status(pair.status):
            updated_pairs.append(pair)
            continue
        collaborator_name = fallback_name or (pair.exam_block.employee_name if pair.exam_block and pair.exam_block.employee_name else None)
        output_name = build_output_filename(collaborator_name, pair.logical_id)
        output_path = os.path.join(pdfs_dir, output_name)
        output_path = ensure_unique_output_path(output_path)
        export_pair_pdf(input_pdf_path, pair, output_path)
        pair.output_file = output_path
        updated_pairs.append(pair)
    return updated_pairs


def export_manual_review_item(
    input_pdf_path: str,
    item: ManualReviewItem,
    pdfs_dir: str,
    fallback_name: Optional[str],
) -> str:
    if item.review_status != "CONFIRMADO_MANUAL":
        raise ValueError("A pendência atual ainda não foi confirmada manualmente.")

    if not item.pair.list_record:
        raise ValueError("A pendência atual não possui lista para exportação.")

    if not item.pair.exam_block:
        raise ValueError("A pendência atual não possui prova confirmada para exportação.")

    collaborator_name = fallback_name or (
        item.pair.exam_block.employee_name
        if item.pair.exam_block and item.pair.exam_block.employee_name
        else None
    )

    output_name = build_output_filename(collaborator_name, item.pair.logical_id)
    output_path = os.path.join(pdfs_dir, output_name)
    output_path = ensure_unique_output_path(output_path)

    export_pair_pdf(input_pdf_path, item.pair, output_path)

    item.generated_output_file = output_path
    item.pair.output_file = output_path

    export_note = "PDF gerado manualmente após confirmação da revisão."
    if export_note not in item.pair.notes:
        item.pair.notes.append(export_note)

    if export_note not in item.notes:
        item.notes.append(export_note)

    return output_path


def export_all_confirmed_manual_review_items(
    input_pdf_path: str,
    items: List[ManualReviewItem],
    pdfs_dir: str,
    fallback_name: Optional[str],
) -> Dict[str, List[str]]:
    exported: List[str] = []
    skipped: List[str] = []
    errors: List[str] = []

    for item in items:
        if item.review_status != "CONFIRMADO_MANUAL":
            skipped.append(f"{item.review_id}: status={item.review_status}")
            continue

        if not item.pair.list_record:
            errors.append(f"{item.review_id}: lista ausente")
            continue

        if not item.pair.exam_block:
            errors.append(f"{item.review_id}: prova confirmada ausente")
            continue

        if item.generated_output_file:
            skipped.append(f"{item.review_id}: já exportado")
            continue

        try:
            output_path = export_manual_review_item(
                input_pdf_path=input_pdf_path,
                item=item,
                pdfs_dir=pdfs_dir,
                fallback_name=fallback_name,
            )
            exported.append(output_path)
        except Exception as exc:
            errors.append(f"{item.review_id}: {exc}")

    return {
        "exported": exported,
        "skipped": skipped,
        "errors": errors,
    }


def get_pair_workflow_bucket(pair: PairRecord) -> str:
    if is_auto_export_status(pair.status):
        return "AUTOMATICO_CONFIAVEL"

    if pair.status == "MANUAL_CONFIRMADO":
        if pair.output_file:
            return "MANUAL_CONFIRMADO_EXPORTADO"
        return "MANUAL_CONFIRMADO_PENDENTE_EXPORTACAO"

    if pair.status == "MANUAL_SEM_CORRESPONDENCIA":
        return "MANUAL_SEM_CORRESPONDENCIA"

    return "PENDENTE_REVISAO_MANUAL"


def summarize_pair_workflow_buckets(pairs: List[PairRecord]) -> Dict[str, int]:
    summary = {
        "AUTOMATICO_CONFIAVEL": 0,
        "MANUAL_CONFIRMADO_EXPORTADO": 0,
        "MANUAL_CONFIRMADO_PENDENTE_EXPORTACAO": 0,
        "MANUAL_SEM_CORRESPONDENCIA": 0,
        "PENDENTE_REVISAO_MANUAL": 0,
    }

    for pair in pairs:
        bucket = get_pair_workflow_bucket(pair)
        summary[bucket] += 1

    return summary


def get_pair_workflow_label(pair: PairRecord) -> str:
    labels = {
        "AUTOMATICO_CONFIAVEL": "Automático confiável",
        "MANUAL_CONFIRMADO_EXPORTADO": "Manual confirmado e exportado",
        "MANUAL_CONFIRMADO_PENDENTE_EXPORTACAO": "Manual confirmado pendente de exportação",
        "MANUAL_SEM_CORRESPONDENCIA": "Manual sem correspondência",
        "PENDENTE_REVISAO_MANUAL": "Pendente de revisão manual",
    }
    bucket = get_pair_workflow_bucket(pair)
    return labels.get(bucket, bucket)


def build_pair_summary_row(pair: PairRecord) -> Dict[str, str]:
    bucket = get_pair_workflow_bucket(pair)

    return {
        "workflow_bucket": bucket,
        "workflow_label": get_pair_workflow_label(pair),
        "logical_id": pair.logical_id,
        "status": pair.status,
        "list_page": str(pair.list_record.page_number) if pair.list_record else "",
        "exam_start_page": str(pair.exam_block.start_page) if pair.exam_block else "",
        "exam_pages": ",".join(map(str, pair.exam_block.pages)) if pair.exam_block else "",
        "output_file": pair.output_file or "",
        "exported": "SIM" if pair.output_file else "NAO",
        "notes": " | ".join(pair.notes) if pair.notes else "",
    }


def build_report_lines(result: ProcessResult) -> List[str]:
    lines: List[str] = []
    summary = summarize_pair_workflow_buckets(result.pairs)
    
    lines.append(f"Sistema: {APP_NAME}")
    lines.append(f"Versão: {APP_VERSION}")
    for line in build_author_report_lines():
        lines.append(line)
    lines.append("")

    lines.append(f"Arquivo de entrada: {result.input_file}")
    lines.append(f"Colaborador identificado: {result.collaborator_name}")
    lines.append(f"Total de páginas: {result.total_pages}")
    lines.append(f"Listas encontradas: {len(result.lists_found)}")
    lines.append(f"Provas encontradas: {len(result.exams_found)}")
    lines.append(f"Pares automáticos confiáveis: {summary['AUTOMATICO_CONFIAVEL']}")
    lines.append(f"Manuais confirmados e exportados: {summary['MANUAL_CONFIRMADO_EXPORTADO']}")
    lines.append(f"Manuais confirmados pendentes de exportação: {summary['MANUAL_CONFIRMADO_PENDENTE_EXPORTACAO']}")
    lines.append(f"Marcados sem correspondência: {summary['MANUAL_SEM_CORRESPONDENCIA']}")
    lines.append(f"Pendentes de revisão manual: {summary['PENDENTE_REVISAO_MANUAL']}")
    lines.append(f"Pasta desta execução: {result.output_dir}")
    lines.append(f"Relatório TXT principal: {result.report_txt_path}")
    lines.append(f"Resumo CSV principal: {result.report_csv_path}")
    lines.append("")

    lines.append("=== RESUMO DO FLUXO ===")
    lines.append(f"- Automático confiável: {summary['AUTOMATICO_CONFIAVEL']}")
    lines.append(f"- Manual confirmado e exportado: {summary['MANUAL_CONFIRMADO_EXPORTADO']}")
    lines.append(f"- Manual confirmado pendente de exportação: {summary['MANUAL_CONFIRMADO_PENDENTE_EXPORTACAO']}")
    lines.append(f"- Manual sem correspondência: {summary['MANUAL_SEM_CORRESPONDENCIA']}")
    lines.append(f"- Pendente de revisão manual: {summary['PENDENTE_REVISAO_MANUAL']}")
    lines.append("")

    grouped_pairs = {
        "AUTOMATICO_CONFIAVEL": [],
        "MANUAL_CONFIRMADO_EXPORTADO": [],
        "MANUAL_CONFIRMADO_PENDENTE_EXPORTACAO": [],
        "MANUAL_SEM_CORRESPONDENCIA": [],
        "PENDENTE_REVISAO_MANUAL": [],
    }

    for pair in result.pairs:
        grouped_pairs[get_pair_workflow_bucket(pair)].append(pair)

    section_titles = [
        ("AUTOMATICO_CONFIAVEL", "=== AUTOMÁTICOS CONFIÁVEIS ==="),
        ("MANUAL_CONFIRMADO_EXPORTADO", "=== MANUAIS CONFIRMADOS E EXPORTADOS ==="),
        ("MANUAL_CONFIRMADO_PENDENTE_EXPORTACAO", "=== MANUAIS CONFIRMADOS PENDENTES DE EXPORTAÇÃO ==="),
        ("MANUAL_SEM_CORRESPONDENCIA", "=== MANUAIS SEM CORRESPONDÊNCIA ==="),
        ("PENDENTE_REVISAO_MANUAL", "=== PENDENTES DE REVISÃO MANUAL ==="),
    ]

    for bucket_key, title in section_titles:
        lines.append(title)
        if grouped_pairs[bucket_key]:
            for pair in grouped_pairs[bucket_key]:
                lines.append(
                    f"- {pair.logical_id} | status={pair.status} | "
                    f"lista={pair.list_record.page_number if pair.list_record else '-'} | "
                    f"prova={pair.exam_block.start_page if pair.exam_block else '-'} | "
                    f"arquivo={pair.output_file or '-'}"
                )
                for note in pair.notes:
                    lines.append(f"    observacao: {note}")
        else:
            lines.append("- Nenhum")
        lines.append("")

    lines.append("=== AVISOS ===")
    if result.warnings:
        for w in result.warnings:
            lines.append(f"- {w}")
    else:
        lines.append("- Nenhum")
    lines.append("")

    lines.append("=== ERROS ===")
    if result.errors:
        for e in result.errors:
            lines.append(f"- {e}")
    else:
        lines.append("- Nenhum")
    lines.append("")

    lines.append("=== PAGINAS CLASSIFICADAS ===")
    for page in result.page_records:
        lines.append(
            f"- Pag {page.page_number}: tipo={page.page_type}, codigo={page.code}, "
            f"anexo={page.annex}, id={page.logical_id}, nome={page.employee_name}, conf={page.confidence:.2f}"
        )
        for note in page.notes[:6]:
            lines.append(f"    nota: {note}")

    return lines


def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"

    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            ws.append(row)

    # Ajuste simples de largura das colunas
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_length:
                max_length = len(value)

        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 60)

    wb.save(xlsx_path)


def write_reports(result: ProcessResult) -> None:
    lines = build_report_lines(result)

    with open(result.report_txt_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    with open(result.report_csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "workflow_bucket",
            "workflow_label",
            "logical_id",
            "status",
            "list_page",
            "exam_start_page",
            "exam_pages",
            "output_file",
            "exported",
            "notes",
        ])
        writer.writeheader()
        for pair in result.pairs:
            writer.writerow(build_pair_summary_row(pair))

    report_xlsx_path = str(Path(result.report_csv_path).with_suffix(".xlsx"))
    csv_to_xlsx(result.report_csv_path, report_xlsx_path)

def process_selected_pdf(file_path: str, progress_callback: Optional[Callable[[str], None]] = None) -> ProcessResult:
    if not file_path:
        raise ValueError("Nenhum arquivo foi selecionado.")

    if progress_callback:
        progress_callback("Iniciando análise das páginas...")
    page_records = analyze_pdf_pages(file_path, progress_callback=progress_callback)

    if progress_callback:
        progress_callback("Coletando listas...")
    lists_found = collect_lists(page_records)

    if progress_callback:
        progress_callback("Montando blocos de prova...")
    exams_found = group_exam_blocks(page_records)

    if progress_callback:
        progress_callback("Definindo colaborador principal...")
    collaborator_name = choose_global_collaborator_name(file_path, exams_found, page_records)

    if progress_callback:
        progress_callback("Preparando pasta de saída...")
    run_output_dir, pdfs_dir, run_id = ensure_output_structure(collaborator_name)

    reports_dir = Path(run_output_dir) / "relatorios"
    reports_dir.mkdir(parents=True, exist_ok=True)

    report_txt_path = str(reports_dir / "relatorio.txt")
    report_csv_path = str(reports_dir / "resumo.csv")

    if progress_callback:
        progress_callback("Pareando listas e provas...")
    pairs, warnings, errors = pair_lists_and_exams(lists_found, exams_found)

    if progress_callback:
        progress_callback("Exportando PDFs finais...")
    pairs = export_ok_pairs(file_path, pairs, pdfs_dir, collaborator_name)

    auto_export_pairs, manual_review_pairs = split_pairs_for_workflow(pairs)
    manual_review_items = build_manual_review_items(manual_review_pairs)

    result = ProcessResult(
        input_file=os.path.basename(file_path),
        input_pdf_path=file_path,
        collaborator_name=collaborator_name or "ColaboradorNaoIdentificado",
        total_pages=len(page_records),
        page_records=page_records,
        lists_found=lists_found,
        exams_found=exams_found,
        pairs=pairs,
        auto_export_pairs=auto_export_pairs,
        manual_review_pairs=manual_review_pairs,
        manual_review_items=manual_review_items,
        warnings=warnings,
        errors=errors,
        output_dir=run_output_dir,
        report_txt_path=report_txt_path,
        report_csv_path=report_csv_path,
    )

    if progress_callback:
        progress_callback("Gravando relatórios...")
    write_reports(result)
    return result


# ============================================================
# Interface gráfica local
# ============================================================

class ColaboradoresApp:
    def __init__(self, root: tk.Tk, license_info: Optional[LicenseValidationResult] = None):
        self.root = root
        self.license_info = license_info
        self.root.title("Separador de Listas e Provas")
        self.root.geometry("900x560")
        self.selected_pdf: Optional[str] = None
        self.last_output_dir: Optional[str] = None
        self.last_result: Optional[ProcessResult] = None
        self.is_busy = False
        self._build_ui()
        self._log("Sistema pronto.")
        self._log("Selecione um PDF e use Pré-análise ou Processar.")

    def build_license_summary_text(self) -> str:
        if not self.license_info:
            return "Licença: informação indisponível"

        if self.license_info.license_type == "perpetual":
            validity_text = "Vitalícia"
        else:
            validity_text = self.license_info.expires_at or "Não informado"

        binding_text = "Sim" if self.license_info.machine_bound else "Não"

        return (
            f"Licenciado para: {self.license_info.customer_name}\n"
            f"Escopo: {self.license_info.customer_type}\n"
            f"Tipo de licença: {self.license_info.license_type}\n"
            f"Validade: {validity_text}\n"
            f"Vinculada à máquina: {binding_text}\n"
            f"License ID: {self.license_info.license_id}"
        )

    def _build_ui(self):
        main = ttk.Frame(self.root, padding=12)
        main.pack(fill="both", expand=True)

        row1 = ttk.Frame(main)
        row1.pack(fill="x", pady=(0, 10))
        ttk.Label(row1, text="PDF selecionado:").pack(side="left")
        self.path_var = tk.StringVar(value="Nenhum arquivo selecionado")
        ttk.Entry(row1, textvariable=self.path_var).pack(side="left", fill="x", expand=True, padx=8)
        self.btn_select_pdf = ttk.Button(row1, text="Selecionar PDF", command=self.select_pdf)
        self.btn_select_pdf.pack(side="left")

        row2 = ttk.Frame(main)
        row2.pack(fill="x", pady=(0, 10))
        self.btn_copy_bruto = ttk.Button(row2, text="Copiar para pasta bruto", command=self.copy_to_bruto)
        self.btn_copy_bruto.pack(side="left")

        self.btn_pre_analysis = ttk.Button(row2, text="Pré-análise", command=self.run_pre_analysis)
        self.btn_pre_analysis.pack(side="left", padx=6)

        self.btn_process = ttk.Button(row2, text="Processar", command=self.run_processing)
        self.btn_process.pack(side="left")

        self.btn_manual_review = ttk.Button(row2, text="Revisão manual", command=self.open_manual_review)
        self.btn_manual_review.pack(side="left", padx=6)

        self.btn_open_output = ttk.Button(row2, text="Abrir saída", command=self.open_last_output)
        self.btn_open_output.pack(side="left", padx=6)

        self.btn_open_base = ttk.Button(row2, text="Abrir pasta base", command=lambda: open_in_file_manager(str(BASE_DIR)))
        self.btn_open_base.pack(side="left")

        info = ttk.Label(
            main,
            text=(
                "Uso local: selecione um PDF e clique em Processar.\n"
                "Os arquivos gerados e os relatórios ficam na pasta da execução dentro de colaboradores_data/processado."
            ),
            justify="left",
        )
        info.pack(anchor="w", pady=(0, 8))

        info_box = ttk.LabelFrame(main, text="Informações do sistema", padding=10)
        info_box.pack(fill="x", pady=(0, 10))

        ttk.Label(
            info_box,
            text=f"{APP_NAME}  |  Versão: {APP_VERSION}",
            justify="left",
        ).pack(anchor="w")

        ttk.Label(
            info_box,
            text=f"Desenvolvedor: {AUTHOR_NAME}",
            justify="left",
        ).pack(anchor="w", pady=(4, 0))

        ttk.Label(
            info_box,
            text=f"E-mail: {AUTHOR_EMAIL}",
            justify="left",
        ).pack(anchor="w", pady=(2, 0))

        ttk.Label(
            info_box,
            text=self.build_license_summary_text(),
            justify="left",
        ).pack(anchor="w", pady=(8, 0))

        self.text = tk.Text(main, wrap="word", height=14)
        self.text.pack(fill="both", expand=True)

        scrollbar = ttk.Scrollbar(self.text, command=self.text.yview)
        scrollbar.pack(side="right", fill="y")
        self.text.configure(yscrollcommand=scrollbar.set)

    def _log(self, message: str):
        self.text.insert("end", message + "\n")
        self.text.see("end")
        self.root.update_idletasks()
            
    def _clear_log(self):
        self.text.delete("1.0", "end")
        self.root.update_idletasks()

    def _start_user_console(self, title: str):
        self._clear_log()
        self._log(title)
        if self.selected_pdf:
            self._log(f"Arquivo: {Path(self.selected_pdf).name}")
        self._log("") 
        
    def _set_busy_state(self, busy: bool):
        self.is_busy = busy
        state = "disabled" if busy else "normal"

        widgets = [
            getattr(self, "btn_select_pdf", None),
            getattr(self, "btn_copy_bruto", None),
            getattr(self, "btn_pre_analysis", None),
            getattr(self, "btn_process", None),
            getattr(self, "btn_manual_review", None),
            getattr(self, "btn_open_output", None),
            getattr(self, "btn_open_base", None),
        ]

        for widget in widgets:
            if widget is not None:
                try:
                    widget.configure(state=state)
                except Exception:
                    pass

        try:
            self.root.configure(cursor="watch" if busy else "")
        except Exception:
            pass

        self.root.update_idletasks()

    def _ui_call(self, callback, *args, **kwargs):
        self.root.after(0, lambda: callback(*args, **kwargs))

    def select_pdf(self):
        if self.is_busy:
            messagebox.showinfo("Aguarde", "Há uma operação em andamento.")
            return
        path = filedialog.askopenfilename(
            title="Selecione o PDF",
            filetypes=[("PDF", "*.pdf"), ("Todos os arquivos", "*.*")],
            initialdir=str(BRUTO_DIR if BRUTO_DIR.exists() else APP_DIR),
        )
        if path:
            clear_preview_caches()
            self.selected_pdf = path
            self.path_var.set(path)
            self._clear_log()
            self._log("PDF selecionado com sucesso.")
            self._log(f"Arquivo: {Path(path).name}")
            self._log(f"Local: {path}")

    def copy_to_bruto(self):
        if self.is_busy:
            messagebox.showinfo("Aguarde", "Há uma operação em andamento.")
            return
        if not self.selected_pdf:
            messagebox.showwarning("Aviso", "Selecione um PDF primeiro.")
            return
        src = Path(self.selected_pdf)
        dst = BRUTO_DIR / src.name
        if src.resolve() == dst.resolve():
            self._log("O arquivo já está na pasta bruto.")
            return
        dst.write_bytes(src.read_bytes())
        self._log(f"Arquivo copiado para: {dst}")


    def run_pre_analysis(self):
        if self.is_busy:
            messagebox.showinfo("Aguarde", "Há uma operação em andamento.")
            return

        if not self.selected_pdf:
            messagebox.showwarning("Aviso", "Selecione um PDF primeiro.")
            return

        selected_pdf = self.selected_pdf
        self._start_user_console("Pré-análise do arquivo selecionado")
        self._log("Lendo estrutura inicial do PDF...")
        self._set_busy_state(True)

        def finish_success(result: PreAnalysisResult):
            self._log("Pré-análise concluída.")
            self._log(f"Arquivo: {result.file_name}")
            self._log(f"Páginas: {result.total_pages}")
            self._log(f"Provavelmente escaneado: {'Sim' if result.is_probably_scanned else 'Não'}")
            self._set_busy_state(False)

        def finish_error(exc: Exception):
            self._log(f"Erro na pré-análise: {exc}")
            self._set_busy_state(False)
            messagebox.showerror("Erro", str(exc))

        def worker():
            try:
                result = pre_analyze_pdf(selected_pdf)
                self._ui_call(finish_success, result)
            except Exception as exc:
                self._ui_call(finish_error, exc)

        threading.Thread(target=worker, daemon=True).start()


    def run_processing(self):
        if self.is_busy:
            messagebox.showinfo("Aguarde", "Há uma operação em andamento.")
            return

        if not self.selected_pdf:
            messagebox.showwarning("Aviso", "Selecione um PDF primeiro.")
            return

        selected_pdf = self.selected_pdf
        self._start_user_console("Processamento em andamento")
        self._set_busy_state(True)

        last_public_progress = {"value": None}

        def safe_log(message: str):
            self._ui_call(self._log, message)

        def friendly_progress(raw_msg: str):
            page_match = re.match(r"Lendo página (\d+) de (\d+)\.\.\.", raw_msg)
            if page_match:
                current = int(page_match.group(1))
                total = int(page_match.group(2))

                if current == 1 or current == total or current % 5 == 0:
                    safe_log(f"Lendo páginas do PDF: {current}/{total}")
                return

            mapping = {
                "Iniciando análise das páginas...": "Analisando o documento...",
                "Coletando listas...": "Separando listas identificadas...",
                "Montando blocos de prova...": "Organizando blocos de prova...",
                "Definindo colaborador principal...": "Identificando o colaborador principal...",
                "Preparando pasta de saída...": "Preparando a pasta de saída...",
                "Pareando listas e provas...": "Relacionando listas e provas...",
                "Exportando PDFs finais...": "Gerando os PDFs automáticos...",
                "Gravando relatórios...": "Atualizando relatórios...",
            }

            public_msg = mapping.get(raw_msg)
            if public_msg and public_msg != last_public_progress["value"]:
                last_public_progress["value"] = public_msg
                safe_log(public_msg)

        def finish_success(result: ProcessResult):
            clear_preview_caches()
            self.last_result = result
            self.last_output_dir = result.output_dir

            total_gerados = len(result.auto_export_pairs)
            total_pendentes = len(result.manual_review_pairs)

            self._log("")
            self._log("Resumo do processamento:")
            self._log(f"Colaborador identificado: {result.collaborator_name}")
            self._log(f"Listas encontradas: {len(result.lists_found)}")
            self._log(f"Provas encontradas: {len(result.exams_found)}")
            self._log(f"PDFs automáticos gerados: {total_gerados}")
            self._log(f"Pendências para revisão manual: {total_pendentes}")
            self._log(f"Relatório principal: {Path(result.report_txt_path).name}")
            self._log(f"Resumo da execução: {Path(result.report_csv_path).with_suffix('.xlsx').name}")
            self._log(f"Pasta de saída: {result.output_dir}")

            self._set_busy_state(False)

            messagebox.showinfo(
                "Concluído",
                f"Processamento concluído.\n\n"
                f"Arquivos gerados automaticamente: {total_gerados}\n"
                f"Pendências para revisão manual: {total_pendentes}"
            )

        def finish_error(exc: Exception, tb_text: str):
            self._log(f"Erro durante o processamento: {exc}")
            self._log(tb_text)
            self._set_busy_state(False)
            messagebox.showerror("Erro", str(exc))

        def worker():
            try:
                result = process_selected_pdf(selected_pdf, progress_callback=friendly_progress)
                self._ui_call(finish_success, result)
            except Exception as exc:
                tb_text = traceback.format_exc()
                self._ui_call(finish_error, exc, tb_text)

        threading.Thread(target=worker, daemon=True).start()


    def open_manual_review(self):
        if self.is_busy:
            messagebox.showinfo("Aguarde", "Finalize a operação atual antes de abrir a revisão manual.")
            return
        if not self.last_result:
            messagebox.showinfo("Revisão manual", "Processe um PDF antes de abrir a revisão manual.")
            return

        if not self.last_result.manual_review_items:
            messagebox.showinfo("Revisão manual", "Não há pendências para revisão manual neste processamento.")
            return

        items = self.last_result.manual_review_items
        review_index = 0
        candidate_index = 0
        list_preview_zoom = 1.35
        candidate_preview_zoom = 1.35
        candidate_preview_page_index = 0
        fit_after_id = None

        win = tk.Toplevel(self.root)
        win.title("Revisão manual assistida")

        screen_w = win.winfo_screenwidth()
        screen_h = win.winfo_screenheight()

        window_w = min(screen_w - 30, max(980, int(screen_w * 0.94)))
        window_h = min(screen_h - 40, max(760, int(screen_h * 0.94)))

        if window_w > screen_w - 40:
            window_w = screen_w - 40

        if window_h > screen_h - 80:
            window_h = screen_h - 80

        pos_x = max(0, (screen_w - window_w) // 2)
        pos_y = max(0, (screen_h - window_h) // 2)

        win.geometry(f"{window_w}x{window_h}+{pos_x}+{pos_y}")
        win.resizable(True, True)
        preview_canvas_height = max(420, int(window_h * 0.58))

        container = ttk.Frame(win)
        container.pack(fill="both", expand=True)

        canvas = tk.Canvas(container, highlightthickness=0)
        v_scroll = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=v_scroll.set)

        v_scroll.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)

        main = ttk.Frame(canvas, padding=12)
        canvas_window = canvas.create_window((0, 0), window=main, anchor="nw")


        def _sync_scrollregion(event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))


        def _resize_inner_frame(event):
            canvas.itemconfigure(canvas_window, width=event.width)


        def _on_mousewheel(event):
            if event.delta:
                canvas.yview_scroll(int(-event.delta / 120), "units")


        def _bind_mousewheel(event=None):
            canvas.bind_all("<MouseWheel>", _on_mousewheel)


        def _unbind_mousewheel(event=None):
            canvas.unbind_all("<MouseWheel>")


        main.bind("<Configure>", _sync_scrollregion)
        canvas.bind("<Configure>", _resize_inner_frame)
        canvas.bind("<Enter>", _bind_mousewheel)
        canvas.bind("<Leave>", _unbind_mousewheel)

        title_var = tk.StringVar()
        summary_var = tk.StringVar()

        progress_line_var = tk.StringVar()

        list_preview_meta_var = tk.StringVar()
        candidate_preview_meta_var = tk.StringVar()
        action_var = tk.StringVar(value="Nenhuma ação manual realizada nesta pendência.")
        code_var = tk.StringVar()
        export_hint_var = tk.StringVar()
        export_button_var = tk.StringVar(value="Gerar todos os PDFs confirmados")

        ttk.Label(
            main,
            textvariable=title_var,
            font=("Segoe UI", 9, "bold"),
        ).pack(anchor="w", pady=(0, 0))

        ttk.Label(
            main,
            textvariable=summary_var,
            justify="left",
        ).pack(anchor="w", pady=(0, 2))
        
        progress_frame = ttk.LabelFrame(main, text="Progresso", padding=3)
        progress_frame.pack(fill="x", pady=(0, 3))

        ttk.Label(
            progress_frame,
            textvariable=progress_line_var,
            justify="left",
        ).pack(anchor="w")

        controls_frame = ttk.Frame(main)
        controls_frame.pack(fill="x", pady=(0, 3))
        controls_frame.columnconfigure(0, weight=2)
        controls_frame.columnconfigure(1, weight=3)
        controls_frame.columnconfigure(2, weight=2)

        navigation_frame = ttk.LabelFrame(controls_frame, text="Navegação", padding=4)
        navigation_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 6))

        actions = ttk.LabelFrame(controls_frame, text="Ações", padding=3)
        actions.grid(row=0, column=1, sticky="nsew", padx=(0, 6))

        code_frame = ttk.LabelFrame(controls_frame, text="Código final", padding=3)
        code_frame.grid(row=0, column=2, sticky="nsew")
        
        preview_container = ttk.Frame(main)
        preview_container.pack(fill="both", expand=True, pady=(0, 2))

        preview_container.columnconfigure(0, weight=1, uniform="preview")
        preview_container.columnconfigure(1, weight=1, uniform="preview")
        preview_container.rowconfigure(0, weight=1)

        list_preview_frame = ttk.LabelFrame(preview_container, text="Prévia da lista", padding=8)
        list_preview_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 6))

        candidate_preview_frame = ttk.LabelFrame(preview_container, text="Prévia da prova candidata", padding=8)
        candidate_preview_frame.grid(row=0, column=1, sticky="nsew", padx=(6, 0))

        list_preview_area = ttk.Frame(list_preview_frame)
        list_preview_area.pack(fill="both", expand=True)

        list_preview_canvas = tk.Canvas(
            list_preview_area,
            highlightthickness=1,
            background="white",
            height=preview_canvas_height,
        )
        list_preview_vscroll = ttk.Scrollbar(
            list_preview_area,
            orient="vertical",
            command=list_preview_canvas.yview,
        )
        list_preview_hscroll = ttk.Scrollbar(
            list_preview_area,
            orient="horizontal",
            command=list_preview_canvas.xview,
        )

        list_preview_canvas.configure(
            yscrollcommand=list_preview_vscroll.set,
            xscrollcommand=list_preview_hscroll.set,
        )

        list_preview_vscroll.pack(side="right", fill="y")
        list_preview_hscroll.pack(side="bottom", fill="x")
        list_preview_canvas.pack(side="left", fill="both", expand=True)
        
        list_preview_actions = ttk.Frame(list_preview_frame)
        list_preview_actions.pack(fill="x", pady=(4, 0))

        candidate_preview_area = ttk.Frame(candidate_preview_frame)
        candidate_preview_area.pack(fill="both", expand=True)

        candidate_preview_canvas = tk.Canvas(
            candidate_preview_area,
            highlightthickness=1,
            background="white",
            height=preview_canvas_height,
        )
        candidate_preview_vscroll = ttk.Scrollbar(
            candidate_preview_area,
            orient="vertical",
            command=candidate_preview_canvas.yview,
        )
        candidate_preview_hscroll = ttk.Scrollbar(
            candidate_preview_area,
            orient="horizontal",
            command=candidate_preview_canvas.xview,
        )

        candidate_preview_canvas.configure(
            yscrollcommand=candidate_preview_vscroll.set,
            xscrollcommand=candidate_preview_hscroll.set,
        )

        candidate_preview_vscroll.pack(side="right", fill="y")
        candidate_preview_hscroll.pack(side="bottom", fill="x")
        candidate_preview_canvas.pack(side="left", fill="both", expand=True)
        
        candidate_preview_actions = ttk.Frame(candidate_preview_frame)
        candidate_preview_actions.pack(fill="x", pady=(4, 0))

        nav_row = ttk.Frame(navigation_frame)
        nav_row.pack(fill="x")

        ttk.Label(
            code_frame,
            text="Corrija o código, se necessário.",
            justify="left",
        ).pack(anchor="w", pady=(0, 2))

        code_entry_row = ttk.Frame(code_frame)
        code_entry_row.pack(fill="x")

        code_entry = ttk.Entry(code_entry_row, textvariable=code_var)
        code_entry.pack(side="left", fill="x", expand=True)

        ttk.Label(
            actions,
            textvariable=action_var,
            justify="left",
        ).pack(anchor="w", pady=(0, 3))

        action_buttons = ttk.Frame(actions)
        action_buttons.pack(fill="x")
        
        ttk.Label(
            actions,
            textvariable=export_hint_var,
            justify="left",
        ).pack(anchor="w", pady=(4, 0))

        footer = ttk.Frame(main)
        footer.pack(fill="x", pady=(2, 0))

        def get_current_item() -> ManualReviewItem:
            return items[review_index]

        def get_current_candidate() -> Optional[ExamBlock]:
            item = get_current_item()
            if not item.candidate_exam_blocks:
                return None
            return item.candidate_exam_blocks[candidate_index]
        
        def get_reserved_candidate_index(item: ManualReviewItem) -> int:
            if item.reserved_exam_start_page is None:
                return 0

            for idx, exam in enumerate(item.candidate_exam_blocks):
                if exam.start_page == item.reserved_exam_start_page:
                    return idx

            return 0
        
        def get_current_candidate_preview_page_number() -> Optional[int]:
            candidate = get_current_candidate()
            if not candidate or not candidate.pages:
                return None

            idx = max(0, min(candidate_preview_page_index, len(candidate.pages) - 1))
            return candidate.pages[idx]

        def bind_preview_pan(canvas_widget: tk.Canvas):
            canvas_widget.bind(
                "<ButtonPress-1>",
                lambda event, c=canvas_widget: c.scan_mark(event.x, event.y),
            )
            canvas_widget.bind(
                "<B1-Motion>",
                lambda event, c=canvas_widget: c.scan_dragto(event.x, event.y, gain=1),
            )
            
        def bind_preview_mouse_controls(canvas_widget: tk.Canvas, preview_kind: str):
            def _on_mousewheel(event):
                ctrl_pressed = (event.state & 0x0004) != 0
                shift_pressed = (event.state & 0x0001) != 0

                if ctrl_pressed:
                    delta = 0.25 if event.delta > 0 else -0.25
                    if preview_kind == "list":
                        change_list_preview_zoom(delta)
                    else:
                        change_candidate_preview_zoom(delta)
                    return "break"

                units = -1 if event.delta > 0 else 1

                if shift_pressed:
                    canvas_widget.xview_scroll(units, "units")
                else:
                    canvas_widget.yview_scroll(units, "units")

                return "break"

            canvas_widget.bind("<Enter>", lambda event, c=canvas_widget: c.focus_set())
            canvas_widget.bind("<MouseWheel>", _on_mousewheel)
            
        bind_preview_pan(list_preview_canvas)
        bind_preview_pan(candidate_preview_canvas)
        bind_preview_mouse_controls(list_preview_canvas, "list")
        bind_preview_mouse_controls(candidate_preview_canvas, "candidate")
        list_preview_canvas.bind("<Configure>", lambda event: schedule_fit_previews())
        candidate_preview_canvas.bind("<Configure>", lambda event: schedule_fit_previews())

        def change_list_preview_zoom(delta: float):
            nonlocal list_preview_zoom
            list_preview_zoom = max(0.8, min(4.0, round(list_preview_zoom + delta, 2)))
            render_current_review()


        def change_candidate_preview_zoom(delta: float):
            nonlocal candidate_preview_zoom
            candidate_preview_zoom = max(0.8, min(4.0, round(candidate_preview_zoom + delta, 2)))
            render_current_review()


        def change_candidate_preview_page(step: int):
            nonlocal candidate_preview_page_index

            candidate = get_current_candidate()
            if not candidate or not candidate.pages:
                return

            new_index = candidate_preview_page_index + step
            if 0 <= new_index < len(candidate.pages):
                candidate_preview_page_index = new_index
                render_current_review()
                
        def reset_list_preview_zoom():
            nonlocal list_preview_zoom
            list_preview_zoom = 1.35
            render_current_review()


        def reset_candidate_preview_zoom():
            nonlocal candidate_preview_zoom
            candidate_preview_zoom = 1.35
            render_current_review()
            
        def fit_list_preview_to_width():
            nonlocal list_preview_zoom

            item = get_current_item()
            if not item.list_record:
                return

            try:
                list_preview_canvas.update_idletasks()
                visible_width = max(
                    200,
                    list_preview_canvas.winfo_width() - list_preview_vscroll.winfo_width() - 12
                )

                if visible_width <= 220:
                    return

                pdf_width, _ = get_pdf_page_size(
                    self.last_result.input_pdf_path,
                    item.list_record.page_number,
                )
                new_zoom = max(0.8, min(4.0, round(visible_width / pdf_width, 2)))

                if abs(list_preview_zoom - new_zoom) >= 0.05:
                    list_preview_zoom = new_zoom
                    render_current_review()
            except Exception:
                pass


        def fit_candidate_preview_to_width():
            nonlocal candidate_preview_zoom

            page_number = get_current_candidate_preview_page_number()
            if not page_number:
                return

            try:
                candidate_preview_canvas.update_idletasks()
                visible_width = max(
                    200,
                    candidate_preview_canvas.winfo_width() - candidate_preview_vscroll.winfo_width() - 12
                )

                if visible_width <= 220:
                    return

                pdf_width, _ = get_pdf_page_size(
                    self.last_result.input_pdf_path,
                    page_number,
                )
                new_zoom = max(0.8, min(4.0, round(visible_width / pdf_width, 2)))

                if abs(candidate_preview_zoom - new_zoom) >= 0.05:
                    candidate_preview_zoom = new_zoom
                    render_current_review()
            except Exception:
                pass
                
        def schedule_fit_previews():
            nonlocal fit_after_id

            def _fit_all():
                nonlocal fit_after_id
                fit_after_id = None
                fit_list_preview_to_width()
                fit_candidate_preview_to_width()

            if fit_after_id is not None:
                try:
                    win.after_cancel(fit_after_id)
                except Exception:
                    pass

            fit_after_id = win.after(120, _fit_all)
        
        def set_preview_canvas_image(canvas_widget: tk.Canvas, pil_image: Optional[Image.Image], empty_text: str):
            canvas_widget.delete("all")

            if pil_image is None:
                canvas_widget.create_text(
                    20,
                    20,
                    anchor="nw",
                    text=empty_text,
                    width=max(canvas_widget.winfo_width() - 40, 200),
                )
                canvas_widget.configure(scrollregion=canvas_widget.bbox("all"))
                canvas_widget.image = None
                return

            photo = ImageTk.PhotoImage(pil_image)
            canvas_widget.create_image(0, 0, anchor="nw", image=photo)
            canvas_widget.image = photo
            canvas_widget.configure(scrollregion=canvas_widget.bbox("all"))
            canvas_widget.xview_moveto(0)
            canvas_widget.yview_moveto(0)

        def update_preview_panels(item: ManualReviewItem, candidate: Optional[ExamBlock]):
            input_pdf_path = self.last_result.input_pdf_path

            list_image = None
            candidate_image = None
            candidate_page_number = get_current_candidate_preview_page_number()

            if item.list_record:
                try:
                    list_image = render_pdf_page_preview(
                        input_pdf_path=input_pdf_path,
                        page_number=item.list_record.page_number,
                        max_size=None,
                        zoom=list_preview_zoom,
                    )
                except Exception:
                    list_image = None

            if candidate and candidate_page_number:
                try:
                    candidate_image = render_pdf_page_preview(
                        input_pdf_path=input_pdf_path,
                        page_number=candidate_page_number,
                        max_size=None,
                        zoom=candidate_preview_zoom,
                    )
                except Exception:
                    candidate_image = None

            set_preview_canvas_image(
                list_preview_canvas,
                list_image,
                "Prévia da lista indisponível."
            )
            set_preview_canvas_image(
                candidate_preview_canvas,
                candidate_image,
                "Prévia da prova candidata indisponível."
            )

            if item.list_record:
                list_preview_meta_var.set(
                    f"PDF pág. {item.list_record.page_number} | Zoom {list_preview_zoom:.2f}x"
                )
            else:
                list_preview_meta_var.set("Sem página de lista.")

            if candidate and candidate_page_number:
                list_page_pos = candidate_preview_page_index + 1
                list_total_pages = len(candidate.pages)
                candidate_preview_meta_var.set(
                    f"Bloco pág. {list_page_pos} de {list_total_pages} | PDF pág. {candidate_page_number} | Zoom {candidate_preview_zoom:.2f}x"
                )
            else:
                candidate_preview_meta_var.set("Sem página de prova disponível.")

        
        def update_manual_review_progress_vars():
            total = len(items)
            confirmed = sum(1 for item in items if item.review_status == "CONFIRMADO_MANUAL")
            no_match = sum(1 for item in items if item.review_status == "SEM_CORRESPONDENCIA")
            exported = sum(1 for item in items if item.generated_output_file)

            decided = confirmed + no_match
            remaining = total - decided

            progress_line_var.set(
                f"Total: {total}   |   Decididas: {decided}   |   Faltam: {remaining}   |   "
                f"Confirmadas: {confirmed}   |   Sem correspondência: {no_match}   |   Exportadas: {exported}"
            )
        
        def sync_item_logical_id(item: ManualReviewItem):
            if item.review_status == "SEM_CORRESPONDENCIA":
                ref_exam = None
            elif item.review_status == "CONFIRMADO_MANUAL":
                ref_exam = item.pair.exam_block
            else:
                ref_exam = item.pair.exam_block or item.suggested_exam_block

            item.pair.logical_id = resolve_manual_review_logical_id(item, ref_exam)


        def apply_typed_code():
            item = get_current_item()
            item.typed_logical_id = normalize_manual_logical_id(code_var.get())
            clear_manual_export_state(item)
            sync_item_logical_id(item)
            render_current_review()

        def clear_typed_code():
            item = get_current_item()
            item.typed_logical_id = ""
            clear_manual_export_state(item)
            sync_item_logical_id(item)
            render_current_review()
            
            
        def clear_manual_export_state(item: ManualReviewItem):
            item.generated_output_file = None
            item.pair.output_file = None

            note_to_remove = "PDF gerado manualmente após confirmação da revisão."
            item.notes = [n for n in item.notes if n != note_to_remove]
            item.pair.notes = [n for n in item.pair.notes if n != note_to_remove]
            
        def export_all_manual_pdfs(close_after_success: bool = False):
            pdfs_dir = str(Path(self.last_result.output_dir) / "pdfs")
            Path(pdfs_dir).mkdir(parents=True, exist_ok=True)

            result_export = export_all_confirmed_manual_review_items(
                input_pdf_path=self.last_result.input_pdf_path,
                items=items,
                pdfs_dir=pdfs_dir,
                fallback_name=self.last_result.collaborator_name,
            )

            write_reports(self.last_result)

            exported_count = len(result_export["exported"])
            error_count = len(result_export["errors"])
            no_match_count = sum(1 for item in items if item.review_status == "SEM_CORRESPONDENCIA")

            self._log("Revisão manual finalizada.")
            self._log(f"PDFs manuais gerados: {exported_count}")
            self._log(f"Listas sem correspondência: {no_match_count}")
            self._log(f"Erros na exportação manual: {error_count}")

            if error_count > 0:
                messagebox.showerror(
                    "Revisão manual",
                    "A classificação foi concluída, mas houve erro(s) na geração automática.\n\n"
                    f"PDFs gerados: {exported_count}\n"
                    f"Sem correspondência: {no_match_count}\n"
                    f"Erros: {error_count}"
                )
                render_current_review()
                return False

            messagebox.showinfo(
                "Revisão manual",
                "Classificação concluída.\n\n"
                f"PDFs gerados: {exported_count}\n"
                f"Sem correspondência: {no_match_count}"
            )

            if close_after_success:
                win.destroy()
            else:
                render_current_review()

            return True 
        

        def update_button_states():
            if review_index <= 0:
                btn_prev_review.state(["disabled"])
            else:
                btn_prev_review.state(["!disabled"])

            if review_index >= len(items) - 1:
                btn_next_review.state(["disabled"])
            else:
                btn_next_review.state(["!disabled"])

            current_item = get_current_item()
            if not current_item.candidate_exam_blocks or candidate_index <= 0:
                btn_prev_candidate.state(["disabled"])
            else:
                btn_prev_candidate.state(["!disabled"])

            if (not current_item.candidate_exam_blocks) or (candidate_index >= len(current_item.candidate_exam_blocks) - 1):
                btn_next_candidate.state(["disabled"])
            else:
                btn_next_candidate.state(["!disabled"])

            if current_item.candidate_exam_blocks:
                btn_accept.state(["!disabled"])
                btn_reject_next.state(["!disabled"])
            else:
                btn_accept.state(["disabled"])
                btn_reject_next.state(["disabled"])

            all_items_classified = all(
                item.review_status in {"CONFIRMADO_MANUAL", "SEM_CORRESPONDENCIA"}
                for item in items
            )

            classified_count = sum(
                1 for item in items
                if item.review_status in {"CONFIRMADO_MANUAL", "SEM_CORRESPONDENCIA"}
            )

            confirmed_pending_export_count = sum(
                1 for item in items
                if item.review_status == "CONFIRMADO_MANUAL"
                and item.pair.list_record
                and item.pair.exam_block
                and not item.generated_output_file
            )

            if all_items_classified and confirmed_pending_export_count > 0:
                btn_export_manual.state(["!disabled"])
                export_button_var.set(f"Gerar {confirmed_pending_export_count} PDF(s) confirmado(s)")
                export_hint_var.set("Classificação concluída. A exportação em lote já está liberada.")
            else:
                btn_export_manual.state(["disabled"])

                if not all_items_classified:
                    remaining_count = len(items) - classified_count
                    export_button_var.set(f"Gerar PDFs ({confirmed_pending_export_count})")
                    export_hint_var.set(
                        f"Finalize a classificação de {remaining_count} lista(s) para liberar a exportação."
                    )
                else:
                    export_button_var.set("Gerar todos os PDFs confirmados")
                    export_hint_var.set("Classificação concluída. Não há PDFs manuais pendentes de exportação.")
                
            current_candidate = get_current_candidate()

            if get_current_item().list_record:
                btn_list_zoom_out.state(["!disabled"])
                btn_list_zoom_in.state(["!disabled"])
                btn_list_fit_width.state(["!disabled"])
                btn_list_zoom_reset.state(["!disabled"])
            else:
                btn_list_zoom_out.state(["disabled"])
                btn_list_zoom_in.state(["disabled"])
                btn_list_fit_width.state(["disabled"])
                btn_list_zoom_reset.state(["disabled"])
                
            if current_candidate and current_candidate.pages:
                btn_candidate_zoom_out.state(["!disabled"])
                btn_candidate_zoom_in.state(["!disabled"])
                btn_candidate_fit_width.state(["!disabled"])
                btn_candidate_zoom_reset.state(["!disabled"])
            else:
                btn_candidate_zoom_out.state(["disabled"])
                btn_candidate_zoom_in.state(["disabled"])
                btn_candidate_fit_width.state(["disabled"])
                btn_candidate_zoom_reset.state(["disabled"])

            if current_candidate and current_candidate.pages and candidate_preview_page_index > 0:
                btn_candidate_page_prev.state(["!disabled"])
            else:
                btn_candidate_page_prev.state(["disabled"])

            if (
                current_candidate
                and current_candidate.pages
                and candidate_preview_page_index < len(current_candidate.pages) - 1
            ):
                btn_candidate_page_next.state(["!disabled"])
            else:
                btn_candidate_page_next.state(["disabled"])

        def render_current_review():
            nonlocal candidate_index, candidate_preview_page_index

            item = get_current_item()
            
            sync_item_logical_id(item)
            code_var.set(item.typed_logical_id or "")

            if item.candidate_exam_blocks:
                candidate_index = max(0, min(candidate_index, len(item.candidate_exam_blocks) - 1))
            else:
                candidate_index = 0

            candidate = get_current_candidate()

            if candidate and candidate.pages:
                candidate_preview_page_index = max(0, min(candidate_preview_page_index, len(candidate.pages) - 1))
            else:
                candidate_preview_page_index = 0

            update_preview_panels(item, candidate)

            title_var.set(
                f"Revisão manual assistida — pendência {review_index + 1} de {len(items)}"
            )

            summary_var.set(
                f"Status: {item.review_status} | Código final: {item.pair.logical_id or '-'}"
            )

            update_manual_review_progress_vars()

            list_page = item.list_record.page_number if item.list_record else "-"
            list_code = item.list_record.logical_id if item.list_record and item.list_record.logical_id else "-"

            list_preview_frame.configure(
                text=f"Lista — pág. {list_page} — código {list_code}"
            )

            if candidate:
                candidate_code = candidate.logical_id or candidate.code or "-"
                candidate_preview_frame.configure(
                    text=f"Prova — {candidate_index + 1}/{len(item.candidate_exam_blocks)} — pág. {candidate.start_page} — código {candidate_code}"
                )
            else:
                candidate_preview_frame.configure(
                    text="Prova — nenhuma candidata disponível"
                )

            if item.review_status == "CONFIRMADO_MANUAL":
                action_var.set(
                    f"Pendência confirmada manualmente com a prova iniciando na página {item.confirmed_exam_start_page}."
                )
            elif item.review_status == "SEM_CORRESPONDENCIA":
                action_var.set("Pendência marcada manualmente como sem correspondência.")
            else:
                action_var.set("Nenhuma ação manual realizada nesta pendência.")

            update_button_states()
            
        def change_review(step: int):
            nonlocal review_index, candidate_index, candidate_preview_page_index
            new_index = review_index + step
            if 0 <= new_index < len(items):
                review_index = new_index
                candidate_index = get_reserved_candidate_index(items[review_index])
                candidate_preview_page_index = 0
                render_current_review()
                schedule_fit_previews()

        def change_candidate(step: int):
            nonlocal candidate_index, candidate_preview_page_index
            item = get_current_item()
            if not item.candidate_exam_blocks:
                return

            new_index = candidate_index + step
            if 0 <= new_index < len(item.candidate_exam_blocks):
                candidate_index = new_index
                candidate_preview_page_index = 0
                render_current_review()
                win.after_idle(fit_candidate_preview_to_width)

        def move_to_next_review_after_decision():
            nonlocal review_index, candidate_index, candidate_preview_page_index

            if review_index < len(items) - 1:
                review_index += 1

            candidate_index = get_reserved_candidate_index(items[review_index])
            candidate_preview_page_index = 0
            render_current_review()
            schedule_fit_previews()
            
        def finish_manual_review_if_ready():
            all_items_classified = all(
                item.review_status in {"CONFIRMADO_MANUAL", "SEM_CORRESPONDENCIA"}
                for item in items
            )

            if not all_items_classified:
                move_to_next_review_after_decision()
                return

            export_all_manual_pdfs(close_after_success=True)

        def accept_current_candidate():
            item = get_current_item()
            candidate = get_current_candidate()

            if not candidate:
                messagebox.showinfo("Revisão manual", "Não existe prova candidata para confirmar nesta pendência.")
                return

            clear_manual_export_state(item)
            reserve_exam_for_review_item(items, item, candidate)
            item.review_status = "CONFIRMADO_MANUAL"
            item.pair.status = "MANUAL_CONFIRMADO"
            sync_item_logical_id(item)

            finish_manual_review_if_ready()

        def reject_and_next_candidate():
            item = get_current_item()

            if not item.candidate_exam_blocks:
                messagebox.showinfo("Revisão manual", "Não existe prova candidata para recusar nesta pendência.")
                return

            if candidate_index >= len(item.candidate_exam_blocks) - 1:
                messagebox.showinfo("Revisão manual", "Esta já é a última prova candidata desta pendência.")
                return

            change_candidate(1)

        def mark_no_match():
            item = get_current_item()

            clear_manual_export_state(item)
            release_exam_for_review_item(items, item, clear_pair_exam=True)
            item.review_status = "SEM_CORRESPONDENCIA"
            item.pair.status = "MANUAL_SEM_CORRESPONDENCIA"
            sync_item_logical_id(item)

            finish_manual_review_if_ready()

        btn_prev_review = ttk.Button(
            nav_row,
            text="Lista anterior",
            command=lambda: change_review(-1),
        )
        btn_prev_review.pack(side="left")

        btn_next_review = ttk.Button(
            nav_row,
            text="Próxima lista",
            command=lambda: change_review(1),
        )
        btn_next_review.pack(side="left", padx=(6, 12))

        btn_prev_candidate = ttk.Button(
            nav_row,
            text="Prova anterior",
            command=lambda: change_candidate(-1),
        )
        btn_prev_candidate.pack(side="left")

        btn_next_candidate = ttk.Button(
            nav_row,
            text="Próxima prova",
            command=lambda: change_candidate(1),
        )
        btn_next_candidate.pack(side="left", padx=6)

        btn_accept = ttk.Button(
            action_buttons,
            text="Aceitar candidata atual",
            command=accept_current_candidate,
        )
        btn_accept.pack(side="left")

        btn_reject_next = ttk.Button(
            action_buttons,
            text="Recusar / próxima prova",
            command=reject_and_next_candidate,
        )
        btn_reject_next.pack(side="left", padx=6)

        btn_no_match = ttk.Button(
            action_buttons,
            text="Marcar sem correspondência",
            command=mark_no_match,
        )
        btn_no_match.pack(side="left")
        
        btn_export_manual = ttk.Button(
            action_buttons,
            textvariable=export_button_var,
            command=export_all_manual_pdfs,
        )
        btn_export_manual.pack(side="left", padx=6)
        
        btn_apply_code = ttk.Button(
            code_entry_row,
            text="Aplicar",
            command=apply_typed_code,
        )
        btn_apply_code.pack(side="left", padx=(6, 0))

        btn_clear_code = ttk.Button(
            code_entry_row,
            text="Limpar",
            command=clear_typed_code,
        )
        btn_clear_code.pack(side="left", padx=6)
        
        btn_list_zoom_out = ttk.Button(
            list_preview_actions,
            text="Zoom -",
            command=lambda: change_list_preview_zoom(-0.25),
        )
        btn_list_zoom_out.pack(side="left")

        btn_list_zoom_in = ttk.Button(
            list_preview_actions,
            text="Zoom +",
            command=lambda: change_list_preview_zoom(0.25),
        )
        btn_list_zoom_in.pack(side="left", padx=6)

        btn_list_fit_width = ttk.Button(
            list_preview_actions,
            text="Ajustar à largura",
            command=fit_list_preview_to_width,
        )
        btn_list_fit_width.pack(side="left", padx=6)

        btn_list_zoom_reset = ttk.Button(
            list_preview_actions,
            text="Resetar zoom",
            command=reset_list_preview_zoom,
        )
        btn_list_zoom_reset.pack(side="left", padx=6)

        ttk.Label(
            list_preview_actions,
            textvariable=list_preview_meta_var,
        ).pack(side="left", padx=12)

        btn_candidate_page_prev = ttk.Button(
            candidate_preview_actions,
            text="Página anterior",
            command=lambda: change_candidate_preview_page(-1),
        )
        btn_candidate_page_prev.pack(side="left")

        btn_candidate_page_next = ttk.Button(
            candidate_preview_actions,
            text="Próxima página",
            command=lambda: change_candidate_preview_page(1),
        )
        btn_candidate_page_next.pack(side="left", padx=6)

        btn_candidate_zoom_out = ttk.Button(
            candidate_preview_actions,
            text="Zoom -",
            command=lambda: change_candidate_preview_zoom(-0.25),
        )
        btn_candidate_zoom_out.pack(side="left", padx=(12, 0))

        btn_candidate_zoom_in = ttk.Button(
            candidate_preview_actions,
            text="Zoom +",
            command=lambda: change_candidate_preview_zoom(0.25),
        )
        btn_candidate_zoom_in.pack(side="left", padx=6)

        btn_candidate_fit_width = ttk.Button(
            candidate_preview_actions,
            text="Ajustar à largura",
            command=fit_candidate_preview_to_width,
        )
        btn_candidate_fit_width.pack(side="left", padx=6)

        btn_candidate_zoom_reset = ttk.Button(
            candidate_preview_actions,
            text="Resetar zoom",
            command=reset_candidate_preview_zoom,
        )
        btn_candidate_zoom_reset.pack(side="left", padx=6)

        ttk.Label(
            candidate_preview_actions,
            textvariable=candidate_preview_meta_var,
        ).pack(side="left", padx=12)

        ttk.Button(footer, text="Fechar", command=win.destroy).pack(side="right")

        render_current_review()
        schedule_fit_previews()
    
    def build_license_summary_text(self) -> str:
        if not self.license_info:
            return "Licença: informação indisponível"

        scope_map = {
            "person": "Pessoa",
            "sector": "Setor",
            "company": "Empresa",
        }

        license_type_map = {
            "temporary": "Temporária",
            "perpetual": "Vitalícia",
        }

        scope_text = scope_map.get(
            self.license_info.customer_type,
            self.license_info.customer_type,
        )

        license_type_text = license_type_map.get(
            self.license_info.license_type,
            self.license_info.license_type,
        )

        if self.license_info.license_type == "perpetual":
            validity_text = "Vitalícia"
        else:
            validity_text = self.license_info.expires_at or "Não informado"

        binding_text = "Sim" if self.license_info.machine_bound else "Não"

        return (
            f"Licenciado para: {self.license_info.customer_name}\n"
            f"Escopo autorizado: {scope_text}\n"
            f"Tipo de licença: {license_type_text}\n"
            f"Validade: {validity_text}\n"
            f"Vinculada à máquina: {binding_text}\n"
            f"Identificador da licença: {self.license_info.license_id}"
        )
    
    
    def open_last_output(self):
        if self.is_busy:
            messagebox.showinfo("Aguarde", "Há uma operação em andamento.")
            return
        if not self.last_output_dir:
            messagebox.showinfo("Saída", "Ainda não há pasta de saída desta sessão.")
            return
        open_in_file_manager(self.last_output_dir)


# ============================================================
# CLI / inicialização
# ============================================================

def open_in_file_manager(path: str) -> None:
    if not os.path.exists(path):
        return
    if sys.platform.startswith("win"):
        os.startfile(path)  # type: ignore[attr-defined]
    elif sys.platform == "darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path])


def install_license_from_dialog(parent: tk.Misc | None = None) -> tuple[bool, str]:
    if filedialog is None:
        return False, "O seletor de arquivos não está disponível nesta máquina."

    selected_path = filedialog.askopenfilename(
        parent=parent,
        title="Selecionar arquivo de licença",
        filetypes=[
            ("Arquivos de licença", "*.json *.lic"),
            ("Arquivos JSON", "*.json"),
            ("Todos os arquivos", "*.*"),
        ],
    )

    if not selected_path:
        return False, ""

    source_path = Path(selected_path)
    target_path = APP_DIR / "license.json"
    backup_path = APP_DIR / "license.json.bak"

    try:
        if target_path.exists():
            if backup_path.exists():
                backup_path.unlink()
            target_path.replace(backup_path)

        shutil.copy2(source_path, target_path)

        validate_license()

        if backup_path.exists():
            backup_path.unlink()

        return (
            True,
            f"Licença instalada com sucesso em:\n{target_path}\n\n"
            "Feche esta janela e abra o programa novamente.",
        )

    except Exception as exc:
        try:
            if target_path.exists():
                target_path.unlink()
        except Exception:
            pass

        try:
            if backup_path.exists():
                backup_path.replace(target_path)
        except Exception:
            pass

        return (
            False,
            "Não foi possível instalar a licença selecionada.\n\n"
            f"Motivo: {exc}"
        )


def show_selectable_license_error(title: str, message: str) -> None:
    if tk is None:
        return

    root = None

    try:
        root = tk.Tk()
        root.title(title)
        root.geometry("780x430")
        root.minsize(700, 360)

        container = tk.Frame(root, padx=12, pady=12)
        container.pack(fill="both", expand=True)

        title_label = tk.Label(
            container,
            text=title,
            font=("Segoe UI", 11, "bold"),
            anchor="w",
            justify="left",
        )
        title_label.pack(fill="x", pady=(0, 8))

        info_label = tk.Label(
            container,
            text="Selecione ou copie as informações abaixo para solicitar a liberação da licença.",
            anchor="w",
            justify="left",
        )
        info_label.pack(fill="x", pady=(0, 8))

        text_widget = tk.Text(
            container,
            wrap="word",
            height=16,
            font=("Consolas", 10),
        )
        text_widget.pack(fill="both", expand=True)
        text_widget.insert("1.0", message)
        text_widget.focus_set()

        button_frame = tk.Frame(container)
        button_frame.pack(fill="x", pady=(10, 0))

        def copy_all() -> None:
            content = text_widget.get("1.0", "end-1c")
            root.clipboard_clear()
            root.clipboard_append(content)
            root.update()

        tk.Button(
            button_frame,
            text="Copiar mensagem",
            command=copy_all,
            width=18,
        ).pack(side="left")
        
        def install_license() -> None:
            ok, feedback = install_license_from_dialog(root)
            if not feedback:
                return

            if ok:
                if messagebox is not None:
                    messagebox.showinfo("Licença", feedback)
                root.destroy()
            else:
                if messagebox is not None:
                    messagebox.showerror("Licença", feedback)

        tk.Button(
            button_frame,
            text="Instalar licença",
            command=install_license,
            width=18,
        ).pack(side="left", padx=(10, 0))

        tk.Button(
            button_frame,
            text="Fechar",
            command=root.destroy,
            width=12,
        ).pack(side="right")

        root.mainloop()

    except Exception:
        if messagebox is not None:
            try:
                messagebox.showerror(title, message)
            except Exception:
                pass
    finally:
        if root is not None:
            try:
                root.destroy()
            except Exception:
                pass


def handle_license_failure(exc: Exception) -> int:
    machine_id = get_current_machine_id() or "indisponível"
    activation_path = None

    try:
        activation_path = save_activation_request(ACTIVATION_REQUEST_PATH)
    except Exception:
        activation_path = None

    exc_text = str(exc).strip()

    if "Arquivo não encontrado" in exc_text and "license.json" in exc_text:
        error_summary = "Nenhum arquivo de licença foi encontrado para esta instalação."
    elif "assinatura" in exc_text.lower():
        error_summary = "O arquivo de licença encontrado é inválido ou foi alterado."
    elif "expirada" in exc_text.lower():
        error_summary = "A licença desta instalação está expirada."
    elif "outro computador" in exc_text.lower():
        error_summary = "A licença informada pertence a outro computador."
    elif "retrocesso" in exc_text.lower():
        error_summary = "Foi detectada inconsistência na data/hora do Windows."
    else:
        error_summary = "Não foi possível validar a licença desta instalação."

    lines = [
        "Licença não validada",
        "====================",
        "",
        error_summary,
        "",
        f"Detalhe técnico: {exc_text}",
        "",
        f"Desenvolvedor responsável: {AUTHOR_NAME}",
        f"E-mail para contato: {AUTHOR_EMAIL}",
        "",
        f"Machine ID: {machine_id}",
    ]

    if activation_path:
        lines.extend(
            [
                "",
                "Foi gerado automaticamente um arquivo de solicitação de ativação em:",
                str(activation_path),
            ]
        )

    lines.extend(
        [
            "",
            "Entre em contato com o desenvolvedor e envie o Machine ID acima",
            "ou o arquivo de solicitação de ativação para emissão da licença correta.",
        ]
    )

    msg = "\n".join(lines)

    print(msg)

    if tk is not None and len(sys.argv) == 1:
        show_selectable_license_error("Licença não validada", msg)

    return 1


def run_cli(pdf_path: str) -> int:
    ensure_base_folders()
    configure_tesseract()
    
    try:
        validate_license()
    except LicenseValidationError as exc:
        return handle_license_failure(exc)

    def cb(msg: str):
        print(msg)

    result = process_selected_pdf(pdf_path, progress_callback=cb)
    total_gerados = sum(1 for p in result.pairs if is_auto_export_status(p.status))
    print()
    print("Processamento concluído com sucesso.")
    print(f"Colaborador: {result.collaborator_name}")
    print(f"Arquivos gerados: {total_gerados}")
    print(f"Pasta de saída: {result.output_dir}")
    print(f"Relatório TXT: {result.report_txt_path}")
    print(f"Resumo CSV: {result.report_csv_path}")
    return 0



def main() -> int:
    ensure_base_folders()
    configure_tesseract()
    
    try:
        license_result = validate_license()
    except LicenseValidationError as exc:
        return handle_license_failure(exc)

    if len(sys.argv) > 1:
        return run_cli(sys.argv[1])

    if tk is None:
        print("Tkinter não está disponível nesta máquina.")
        return 1

    root = tk.Tk()
    style = ttk.Style()
    try:
        style.theme_use("vista")
    except Exception:
        pass
    app = ColaboradoresApp(root, license_info=license_result)
    root.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
